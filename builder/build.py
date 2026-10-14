"""合格一貫カリキュラム マスターテンプレート（.xlsm）を組み立てる。

    python3 builder/build.py <元の試作ファイル.xlsm> <出力先.xlsm>

元の試作ファイルを土台にして見た目（書式・色・罫線）を引き継ぎ、
  - 管理表に「高2の1月」列を追加（17期間：高2の1月〜直前Ⅱ期）
  - 生徒情報（入力）・志望校マスタ・講座マスタ・教材マスタ・カリキュラムDB・共テDB・使い方 を追加
  - VBA（vba/ 以下のソース）から vbaProject.bin を生成して .xlsm にする
"""
import os
import sys
import tempfile
from copy import copy

import openpyxl
from openpyxl.formatting.rule import DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.cell_range import CellRange
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import data as D  # noqa: E402
from package import make_xlsm  # noqa: E402
from vbaproject import build_vba_project  # noqa: E402

FONT = "游ゴシック"
NAVY = "17358F"
INK = "1D2433"
GRAY = "6B7488"
PALE = "EEF2FB"
YELLOW_IN = "FFF6CC"   # 入力セル（黄色）
GOLD = "FFD678"

PLAN = "合格一貫カリキュラム"
FIRST, BLOCK, NSUB, NPER = 8, 6, 12, D.NPER
LAST = FIRST + NSUB * BLOCK - 1          # 79
COL_PROG, COL_CHECK, COL_KY = 19, 21, 22  # S, U, V
KY_COLS = range(23, 27)                   # W..Z
SET_FIRST = 13
MAT_FIRST, MAT_N = 20, 10

THIN = Side(style="thin", color="BEC3CD")
MED = Side(style="medium", color=NAVY)


def font(sz=10, b=False, color=INK, **kw):
    return Font(name=FONT, sz=sz, b=b, color=color, **kw)


def fill(rgb):
    return PatternFill("solid", fgColor=rgb)


def tint(rgb, t):
    """rgb を白に向けて t（0〜1）だけ薄める"""
    r, g, b = (int(rgb[i:i + 2], 16) for i in (0, 2, 4))
    f = lambda x: round(x + (255 - x) * t)  # noqa: E731
    return "%02X%02X%02X" % (f(r), f(g), f(b))


def box(ws, rng, side=THIN):
    for row in ws[rng]:
        for c in row:
            c.border = Border(left=side, right=side, top=side, bottom=side)


# ============================================================ シートのコピー（行・列の付け替えつき）
def copy_sheet(src, dst, rmap=lambda r: r, cmap=lambda c: c, keep_value=lambda r, c: True,
               keep_merge=lambda m: True):
    merges = []
    for m in src.merged_cells.ranges:
        if not keep_merge(m):
            continue
        r1, r2 = rmap(m.min_row), rmap(m.max_row)
        c1 = 2 if m.min_col == 2 and cmap(2) != 2 else cmap(m.min_col)
        c2 = cmap(m.max_col)
        if None in (r1, r2, c1, c2):
            continue
        merges.append((r1, c1, r2, c2))
    for r1, c1, r2, c2 in merges:
        dst.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
    merged_tl = {(r1, c1) for r1, c1, _, _ in merges}
    merged_any = set()
    for r1, c1, r2, c2 in merges:
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                merged_any.add((r, c))
    for row in src.iter_rows():
        for c in row:
            nr, nc = rmap(c.row), cmap(c.column)
            if nr is None or nc is None:
                continue
            if c.column == 2 and cmap(2) != 2:
                targets = [2, nc]          # 新しいB列は元のB列の書式を引き継ぐ
            else:
                targets = [nc]
            for tc in targets:
                d = dst.cell(nr, tc)
                d._style = copy(c._style)
                if (nr, tc) in merged_any and (nr, tc) not in merged_tl:
                    continue
                if c.value is not None and keep_value(c.row, c.column):
                    vc = 2 if (c.column == 2 and (nr, 2) in merged_tl) else nc
                    if tc == vc:
                        dst.cell(nr, vc).value = c.value
    for r, dim in src.row_dimensions.items():
        nr = rmap(r)
        if nr is not None and dim.height is not None:
            dst.row_dimensions[nr].height = dim.height
    for letter, dim in src.column_dimensions.items():
        if dim.width is None:
            continue
        idx = openpyxl.utils.column_index_from_string(letter)
        for i in range(dim.min or idx, (dim.max or idx) + 1):
            nc = cmap(i)
            if nc is not None:
                dst.column_dimensions[L(nc)].width = dim.width
    dst.sheet_view.showGridLines = src.sheet_view.showGridLines
    dst.sheet_view.zoomScale = src.sheet_view.zoomScale
    for a in ("orientation", "paperSize", "scale", "fitToWidth", "fitToHeight"):
        setattr(dst.page_setup, a, getattr(src.page_setup, a))
    dst.page_margins = copy(src.page_margins)
    dst.print_options = copy(src.print_options)
    dst.sheet_properties.pageSetUpPr = copy(src.sheet_properties.pageSetUpPr)
    dst.sheet_properties.tabColor = copy(src.sheet_properties.tabColor)


def copy_cf(src, dst, rmap=lambda r: r, cmap=lambda c: c, keep=lambda rng: True):
    for cf in src.conditional_formatting:
        for rng in cf.sqref.ranges:
            if not keep(rng):
                continue
            r1, r2, c1, c2 = rmap(rng.min_row), rmap(rng.max_row), cmap(rng.min_col), cmap(rng.max_col)
            ref = f"{L(c1)}{r1}:{L(c2)}{r2}"
            for rule in cf.rules:
                dst.conditional_formatting.add(ref, copy(rule))


# ============================================================ Python 版の反映ロジック（初期表示用・VBA と同じ規則）
def lvl_num(x, th_hard=60, th_std=50):
    return 3 if x >= th_hard else 2 if x >= th_std else 1


def judge(inp, univ):
    """生徒情報の数式と同じ判定"""
    u = univ.get(inp["school1"])
    kubun = u[0] if u else "国公立"
    bunri = inp["bunri"] or "文系"
    system = ("私大" if kubun == "私立" else "国立") + bunri
    if inp.get("target_manual") not in (None, ""):
        target = inp["target_manual"]
    elif u:
        b, r = u[1], u[2]
        target = (r if r is not None else b) if bunri == "理系" else (b if b is not None else r)
    else:
        target = None
    cur = inp.get("current")
    pre = lvl_num(cur) if cur not in (None, "") else (lvl_num(target) if target is not None else 2)
    tgt = lvl_num(target) if target is not None else pre
    post = max(pre, min(tgt, pre + 1))
    return dict(kubun=kubun, bunri=bunri, system=system, target=target,
                pre=D.LEVELS[pre - 1], post=D.LEVELS[post - 1])


def recommend_subjects(j, inp):
    kok = j["system"].startswith("国立")
    bunri, s1, s2, sci = j["bunri"], inp["social1"], inp["social2"], inp["science"]
    ky_use = inp["kyote_use"] == "する"
    out = {}
    for subj in D.SUBJECT_NAMES:
        if subj == "英語":
            e = True
        elif subj == "数学":
            e = bunri == "理系" or kok or s1 == "数学"
        elif subj in ("現代文", "古文"):
            e = bunri == "文系"
        elif subj == "漢文":
            e = bunri == "文系" and kok
        elif subj in ("日本史", "世界史", "地理"):
            e = bunri == "文系" and s1 == subj
        elif subj in ("物理", "化学", "生物"):
            e = bunri == "理系" and subj in sci
        else:
            e = False
        if kok:
            if subj in ("英語", "数学", "現代文", "古文", "漢文", "公共・情報"):
                k = True
            elif subj in ("日本史", "世界史", "地理"):
                k = s1 == subj or s2 == subj
            else:
                k = subj in sci
        else:
            k = e and ky_use
        out[subj] = ("○" if e else "", "○" if k else "")
    return out


def reflect(ws_plan, j, inp, subjects, db_index, school):
    """管理表に書き込む（VBA の BuildCurriculum と同じ）"""
    for m in list(ws_plan.merged_cells.ranges):
        if m.min_row >= FIRST and m.max_row <= LAST and m.min_col >= 2 and m.max_col <= 1 + NPER:
            ws_plan.unmerge_cells(str(m))
    for r in range(FIRST, LAST + 1):
        for c in list(range(2, 2 + NPER)) + [COL_CHECK, COL_KY] + list(KY_COLS):
            ws_plan.cell(r, c).value = None
    for s, subj in enumerate(D.SUBJECT_NAMES):
        r = FIRST + s * BLOCK
        ko, ky = subjects[subj]
        if ko == "○":
            for off, kind in ((0, "講座"), (1, "回"), (2, "塾"), (4, "家")):
                vals = []
                for c in range(NPER):
                    lv = j["pre"] if c < D.SPLIT else j["post"]
                    row = db_index.get((j["system"], subj, lv, kind))
                    vals.append(row[c] if row else "")
                c = 0
                while c < NPER:
                    v = vals[c]
                    if v in ("", "←"):
                        c += 1
                        continue
                    e = c
                    while e + 1 < NPER and vals[e + 1] == "←":
                        e += 1
                    ws_plan.cell(r + off, 2 + c).value = v.replace("{志望校}", school)
                    if e > c:
                        ws_plan.merge_cells(start_row=r + off, start_column=2 + c, end_row=r + off, end_column=2 + e)
                    if off in (2, 4) and not v.startswith("（"):
                        for k in range(c, e + 1):
                            ws_plan.cell(r + off + 1, 2 + k).value = "☐"
                    c = e + 1
        if ky == "○":
            k = 0
            soc = inp["social1"] + "/" + inp["social2"]
            for (subj2, name, m9, m10, m11, m12, tag) in D.KYOTE:
                if subj2 != subj:
                    continue
                ok = (tag in ("全", "情報") or (tag in ("文系", "理系") and tag == j["bunri"])
                      or (tag == "公共政経" and "政経" in soc) or (tag == "公共倫理" and "倫理" in soc))
                if ok and k < BLOCK:
                    for i, v in enumerate((name, m9, m10, m11, m12)):
                        ws_plan.cell(r + k, COL_KY + i).value = v or None
                    k += 1


def materials_for(ws_plan, subjects):
    keys = [m[0] for m in D.MATERIALS]
    names = []
    any_ky = any(v[1] == "○" for v in subjects.values())
    for s, subj in enumerate(D.SUBJECT_NAMES):
        if subjects[subj][0] != "○":
            continue
        for c in range(2, 2 + NPER):
            txt = ws_plan.cell(FIRST + s * BLOCK + 4, c).value or ""
            if not txt:
                continue
            for key in keys:
                if key in ("共通テスト過去問", "過去問"):
                    continue
                if key in txt and key not in names and len(names) < MAT_N - 2:
                    names.append(key)
    if any_ky:
        names.append("共通テスト過去問")
    names.append("過去問")
    return [next(m for m in D.MATERIALS if m[0] == n) for n in names]


# ============================================================ メイン
def build(src_path, out_path, sample=None):
    sample = sample or dict(name="", highschool="", grade="高2", tantou="",
                            school1="一橋大学", school2="", school3="", bunri="文系",
                            current=57.5, target_manual=None, social1="日本史", social2="公共・政経",
                            science="化学基礎・生物基礎", kyote_use="しない", auto="する")
    wb = openpyxl.load_workbook(src_path)
    old_plan, old_mock, old_rep, old_set = wb[PLAN], wb["模試・自学"], wb["成績表"], wb["設定"]
    for ws in (old_plan, old_mock, old_rep, old_set):
        ws.title = "_old_" + ws.title
    univ = {u[0]: (u[1], u[2], u[3]) for u in D.UNIVERSITIES}

    # ---------------------------------------------------------------- シート作成（並び順）
    ws_how = wb.create_sheet("使い方")
    ws_in = wb.create_sheet("生徒情報")
    ws_plan = wb.create_sheet(PLAN)
    ws_mock = wb.create_sheet("模試・自学")
    ws_rep = wb.create_sheet("成績表")
    ws_set = wb.create_sheet("設定")
    ws_univ = wb.create_sheet("志望校マスタ")
    ws_mat = wb.create_sheet("教材マスタ")
    ws_course = wb.create_sheet("講座マスタ")
    ws_db = wb.create_sheet("カリキュラムDB")
    ws_ky = wb.create_sheet("共テDB")
    codenames = {ws_how: "shHowTo", ws_in: "shInput", ws_plan: "shPlan", ws_mock: "shMock", ws_rep: "shReport",
                 ws_set: "shSetting", ws_univ: "shUniv", ws_mat: "shMaterial", ws_course: "shCourse",
                 ws_db: "shDB", ws_ky: "shKyote"}
    for ws, cn in codenames.items():
        ws.sheet_properties.codeName = cn

    # 名前の定義
    def name(n, ref):
        if n in wb.defined_names:
            del wb.defined_names[n]
        wb.defined_names[n] = DefinedName(n, attr_text=ref)

    # ================================================================ 設定
    build_settings(ws_set, old_set, name)
    # ================================================================ 生徒情報
    build_input(ws_in, name, sample)
    # ================================================================ マスタ類
    build_masters(ws_univ, ws_mat, ws_course, ws_db, ws_ky, name)
    # ================================================================ 管理表
    build_plan(ws_plan, old_plan)
    # ================================================================ 模試・自学 / 成績表
    build_mock(ws_mock, old_mock)
    build_report(ws_rep, old_rep)
    build_howto(ws_how)

    # 初期表示：サンプル入力でおすすめを反映しておく
    j = judge(sample, univ)
    subjects = recommend_subjects(j, sample)
    for i, subj in enumerate(D.SUBJECT_NAMES):
        ws_set.cell(SET_FIRST + i, 2).value = subjects[subj][0] or None
        ws_set.cell(SET_FIRST + i, 3).value = subjects[subj][1] or None
    db_index = {(r[0], r[1], r[2], r[3]): r[4] for r in D.build_db()}
    reflect(ws_plan, j, sample, subjects, db_index, sample["school1"] or "志望校")
    mats = materials_for(ws_plan, subjects)
    for i, m in enumerate(mats):
        r = MAT_FIRST + i
        ws_mock.cell(r, 1).value, ws_mock.cell(r, 2).value = m[1], m[2]
        ws_mock.cell(r, 3).value = m[3] if m[3] != "" else None
        ws_mock.cell(r, 6).value = m[4]
    for i, subj in enumerate(D.SUBJECT_NAMES):     # 表示/非表示（科目を反映と同じ）
        r = FIRST + i * BLOCK
        ko, ky = subjects[subj]
        show = bool(ko or ky)
        for k in range(BLOCK):
            hide = not show or (show and not ko and k >= 2 and not ws_plan.cell(r + k, COL_KY).value)
            ws_plan.row_dimensions[r + k].hidden = hide

    for ws in (old_plan, old_mock, old_rep, old_set):
        wb.remove(ws)
    # 既存の名前（期間終了日など）は設定シートを指すよう作り直し済み。古い印刷範囲は削除
    for n in list(wb.defined_names):
        if n.startswith("_xlnm"):
            del wb.defined_names[n]
    ws_plan.print_area = f"A1:Z{LAST + 2}"
    ws_rep.print_area = "A1:L60"
    wb.code_name = "ThisWorkbook"
    wb.active = wb.sheetnames.index("生徒情報")
    for ws in wb.worksheets:
        ws.sheet_view.tabSelected = ws.title == "生徒情報"

    tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False).name
    wb.save(tmp)

    vba_dir = os.path.join(HERE, "vba")
    rd = lambda f: open(os.path.join(vba_dir, f), encoding="utf-8").read()  # noqa: E731
    modules = [("ThisWorkbook", "workbook", "Option Explicit\n")]
    for ws, cn in codenames.items():
        code = {"shPlan": rd("shPlan.cls"), "shInput": rd("shInput.cls")}.get(cn, "Option Explicit\n")
        modules.append((cn, "sheet", code))
    modules += [("modCurriculum", "std", rd("modCurriculum.bas")), ("modReport", "std", rd("modReport.bas"))]
    vba = build_vba_project(os.path.join(HERE, "vba_template.bin"), modules)

    btn = dict(fill=NAVY)
    buttons = {
        "生徒情報": [
            dict(col=4, row=3, col2=7, row2=6, text="おすすめカリキュラムを反映", macro="ApplyRecommended", size=1200, **btn),
            dict(col=4, row=7, col2=7, row2=9, text="成績表を印刷", macro="PrintReport", fill="3462B0"),
            dict(col=4, row=10, col2=7, row2=12, text="管理表（A3）を印刷", macro="PrintPlan", fill="3462B0"),
        ],
        PLAN: [
            dict(col=27, row=0, col2=30, row2=3, text="成績表を印刷", macro="PrintReport", **btn),
            dict(col=27, row=6, col2=30, row2=8, text="科目を反映", macro="ApplySubjects", fill="3462B0"),
            dict(col=27, row=12, col2=30, row2=15, text="おすすめカリキュラムを反映", macro="ApplyRecommended", fill="E88A12", size=1000),
        ],
        "設定": [dict(col=7, row=11, col2=10, row2=13, text="科目を反映", macro="ApplySubjects", **btn)],
    }
    make_xlsm(tmp, out_path, vba, buttons)
    os.unlink(tmp)
    return out_path


# ============================================================ 設定シート
def build_settings(ws, old, name):
    copy_sheet(old, ws, keep_value=lambda r, c: r in (1, 2, 8, 9, 11, 12) or (r >= 13 and c == 1))
    lab, note = old["A1"]._style, old["C1"]._style
    ws["B1"] = 2027
    ws["C1"] = "← 生徒の高3の年度（例：2027年度の高3なら 2027。開始は高2の1月＝2027年1月）"
    ws["B2"] = "=TODAY()"
    hdr_style, date_style = old["B4"]._style, old["B5"]._style
    hl_style = None
    ws["A4"], ws["A5"] = "期間", "終了日"
    ws["A4"]._style, ws["A5"]._style = old["A4"]._style, old["A5"]._style
    for i, (p, _, f) in enumerate(D.PERIODS):
        c = ws.cell(4, 2 + i, f"={PLAN}!{L(2 + i)}7")
        c._style = copy(hdr_style)
        d = ws.cell(5, 2 + i, f.replace("年度", "$B$1"))
        d._style = copy(date_style)
        d.number_format = "m/d"
        ws.column_dimensions[L(2 + i)].width = 9.6
    ws["A6"] = ("※ 終了日は仮の値（要確認）。塾の年間予定に合わせて書き換えてください。"
                "終了日を過ぎても受講済み(☑)でない回・☐は、管理表で赤く表示されます。")
    ws["A8"], ws["B8"] = "偏差値の基準", 40
    # 受験科目
    ws["A11"] = "受験科目（「おすすめカリキュラムを反映」で E・F 列の○が B・C 列にコピーされます。手で変えたら「科目を反映」ボタン）"
    for col, txt in ((5, "おすすめ：個別"), (6, "おすすめ：共テ")):
        c = ws.cell(12, col, txt)
        c._style = copy(ws["C12"]._style)
        ws.column_dimensions[L(col)].width = 13
    sysx, bun, s1, s2, sci, kyu = "kSystem", "kBunri", "kSocial1", "kSocial2", "kScience", "kKyoteUse"
    kok = f'LEFT({sysx},2)="国立"'
    for i, subj in enumerate(D.SUBJECT_NAMES):
        r = SET_FIRST + i
        A = f"$A{r}"
        if subj == "英語":
            e = '"○"'
        elif subj == "数学":
            e = f'IF(OR({bun}="理系",{kok},{s1}="数学"),"○","")'
        elif subj in ("現代文", "古文"):
            e = f'IF({bun}="文系","○","")'
        elif subj == "漢文":
            e = f'IF(AND({bun}="文系",{kok}),"○","")'
        elif subj in ("日本史", "世界史", "地理"):
            e = f'IF(AND({bun}="文系",{s1}={A}),"○","")'
        elif subj in ("物理", "化学", "生物"):
            e = f'IF(AND({bun}="理系",ISNUMBER(SEARCH({A},{sci}))),"○","")'
        else:
            e = '""'
        if subj in ("英語", "数学", "現代文", "古文", "漢文", "公共・情報"):
            k_kok = '"○"'
        elif subj in ("日本史", "世界史", "地理"):
            k_kok = f'IF(OR({s1}={A},{s2}={A}),"○","")'
        else:
            k_kok = f'IF(ISNUMBER(SEARCH({A},{sci})),"○","")'
        ws.cell(r, 5).value = "=" + e
        ws.cell(r, 6).value = f'=IF({kok},{k_kok},IF({kyu}="する",E{r},""))'
        for col in (5, 6):
            ws.cell(r, col)._style = copy(ws.cell(r, 4)._style)
            ws.cell(r, col).font = font(10, True, GRAY)
            ws.cell(r, col).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(r, 4).value = f'=IF(OR(B{r}="○",C{r}="○"),"表示","非表示")'
    ws["A25"] = "※ B・C列（○）は手で変更できます。変更後は「科目を反映」ボタンで管理表の表示/非表示を更新。"
    dv = DataValidation(type="list", formula1='"○"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"B{SET_FIRST}:C{SET_FIRST + NSUB - 1}")
    copy_cf(old, ws)

    # レベル判定の基準・系統メモ
    r0 = 27
    ws.cell(r0, 1, "レベル判定の基準（偏差値）")._style = copy(old["A11"]._style)
    rows = [("難関レベル", 60, "← この偏差値以上を『難関』とする（例：旧帝大・一橋・早慶上智など）", "難関閾値"),
            ("標準レベル", 50, "← この偏差値以上を『標準』、未満を『基礎』とする", "標準閾値")]
    for i, (a, b, c, n) in enumerate(rows):
        r = r0 + 1 + i
        ws.cell(r, 1, a)._style = copy(lab)
        ws.cell(r, 2, b)._style = copy(old["B8"]._style)
        ws.cell(r, 3, c)._style = copy(note)
        name(n, f"'設定'!$B${r}")
    ws.cell(r0 + 3, 1, "※ 前期（高2の1月〜夏期）は『現在の偏差値』、後期（9月〜直前）は『志望校の目標偏差値』でレベルを決めます。"
                       "ただし後期は前期の1段階上まで（差が大きいときは講師と相談）。")._style = copy(note)
    r1 = r0 + 5
    ws.cell(r1, 1, "系統ごとの入試の目安（生徒情報シートに表示）")._style = copy(old["A11"]._style)
    for i, sysname in enumerate(D.SYSTEMS):
        r = r1 + 1 + i
        ws.cell(r, 1, sysname)._style = copy(lab)
        ws.cell(r, 2, D.SCHOOL_NOTES[sysname])._style = copy(note)
    name("系統メモ", f"'設定'!$A${r1 + 1}:$B${r1 + 4}")
    name("年度", "'設定'!$B$1")
    name("基準日", "'設定'!$B$2")
    name("期間終了日", f"'設定'!$B$5:${L(1 + NPER)}$5")
    name("偏差値基準", "'設定'!$B$8")
    ws.column_dimensions["A"].width = 14
    ws.sheet_properties.tabColor = "828CA0"


# ============================================================ 生徒情報シート
def build_input(ws, name, s):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 58
    ws.column_dimensions["D"].width = 3
    for c in "EFGH":
        ws.column_dimensions[c].width = 9
    ws.column_dimensions["I"].width = 3
    ws.column_dimensions["J"].width = 14
    ws.sheet_properties.tabColor = "E88A12"
    ws.merge_cells("A1:C1")
    ws["A1"] = "生徒情報・志望校の入力"
    ws["A1"].font, ws["A1"].fill = font(15, True, "FFFFFF"), fill(NAVY)
    ws["A1"].alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 30
    ws.merge_cells("A2:C2")
    ws["A2"] = ("黄色のセルに入力します。志望校・文理・偏差値・選択科目を入れると、"
                "おすすめカリキュラムが管理表に自動で反映されます（マクロを有効にして開いてください）。")
    ws["A2"].font = font(9, color=GRAY)
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[2].height = 30

    def section(r, text):
        ws.merge_cells(f"A{r}:C{r}")
        c = ws.cell(r, 1, text)
        c.font, c.fill = font(11, True, "FFFFFF"), fill("1F4E9E")
        c.alignment = Alignment(vertical="center", indent=1)
        ws.row_dimensions[r].height = 20

    def row(r, label, value=None, note="", inp=True, fmt=None):
        a = ws.cell(r, 1, label)
        a.font, a.fill = font(10, True, NAVY), fill(PALE)
        a.alignment = Alignment(vertical="center", indent=1)
        b = ws.cell(r, 2, value)
        b.font = font(11, True, INK if inp else NAVY)
        b.fill = fill(YELLOW_IN) if inp else fill("FFFFFF")
        b.alignment = Alignment(horizontal="center", vertical="center", shrink_to_fit=True)
        if fmt:
            b.number_format = fmt
        c = ws.cell(r, 3, note)
        c.font = font(8.5, color=GRAY)
        c.alignment = Alignment(vertical="center", wrap_text=True)
        for x in (a, b, c):
            x.border = Border(bottom=THIN, top=THIN)
        ws.row_dimensions[r].height = 21

    section(4, "■ 生徒")
    row(5, "氏名", s["name"] or None, "成績表・管理表に表示されます")
    row(6, "高校", s["highschool"] or None)
    row(7, "学年", s["grade"], "高2の1月スタートが標準（高3・既卒も可）")
    row(8, "担当", s["tantou"] or None)
    section(10, "■ 志望校と成績（ここを変えると自動で反映）")
    row(11, "第1志望校", s["school1"] or None, "リストから選択（志望校マスタ）。区分（国公立/私立）と目標偏差値はここで決まります")
    row(12, "第2志望校", s["school2"] or None, "表示用（判定には使いません）")
    row(13, "第3志望校", s["school3"] or None, "表示用（判定には使いません）")
    row(14, "文理", s["bunri"], "文系 / 理系")
    row(15, "現在の偏差値", s["current"], "直近の模試の総合偏差値（全統模試など）。前期のレベルを決めます", fmt="0.0")
    row(16, "目標偏差値（手入力・任意）", s["target_manual"], "空欄なら志望校マスタの値を使います", fmt="0.0")
    row(17, "地歴公民（第1選択）", s["social1"], "文系：個別試験・共テで使う科目。理系：共テの社会。私大文系で数学受験なら『数学』")
    row(18, "地歴公民（第2・共テのみ）", s["social2"], "国公立文系で共テに2科目必要な場合など。不要なら『なし』")
    row(19, "理科", s["science"], "文系は基礎2科目、理系は2科目を選択")
    row(20, "私大の共通テスト利用", s["kyote_use"], "私立志望で共通テスト利用入試も受けるなら『する』")
    row(21, "入力したら自動で反映", s["auto"], "『しない』にすると、ボタンを押したときだけ反映します")

    section(23, "■ 自動判定（数式・入力不要）")
    J = "$J"
    calc = [
        (24, "区分", f'=IF(B11="","国公立",IFERROR(INDEX(志望校区分,MATCH(B11,志望校リスト,0)),"国公立"))', ""),
        (25, "系統", '=IF(B24="私立","私大","国立")&J14', "国立文系 / 国立理系 / 私大文系 / 私大理系"),
        (26, "志望校の目標偏差値", '=IF(B16<>"",B16,IF(J14="理系",IF(K26<>"",K26,J26),IF(J26<>"",J26,K26)))',
         "志望校マスタの値（学部・方式で違うので目安）"),
        (27, "目標との差（目標−現在）", '=IF(OR(B26="",B15=""),"",B26-B15)', "プラスなら、その分だけ伸ばす必要があります"),
        (28, "前期レベル（高2の1月〜夏期）", '=CHOOSE(J28,"基礎","標準","難関")', "現在の偏差値で判定"),
        (29, "後期レベル（9月〜直前）", '=CHOOSE(J30,"基礎","標準","難関")', "志望校の目標偏差値で判定（前期の1段階上まで）"),
        (30, "判定コメント",
         '=IF(J29>J28+1,"目標との差が大きいため、後期は1段階上のレベルまでにしています（講師と相談して調整）",'
         'IF(J29>J28,"後期でレベルを1段階上げて、志望校レベルに合わせます","前期・後期とも同じレベルで仕上げます"))', ""),
        (31, "入試の目安", '=IFERROR(VLOOKUP(B25,系統メモ,2,FALSE),"")', ""),
        (32, "おすすめパターン", '=B25&"／前期："&B28&"／後期："&B29', "カリキュラムDB のこの組み合わせを管理表に反映します"),
    ]
    for r, label, f, note in calc:
        row(r, label, f, note, inp=False, fmt="0.0" if r in (26, 27) else None)
    for r in (30, 31):
        ws.merge_cells(f"B{r}:C{r}")
        ws.cell(r, 2).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.cell(r, 2).font = font(9, True, NAVY)
        ws.row_dimensions[r].height = 30
    # 補助セル（J〜K列）
    ws["J22"] = "補助（数式）"
    ws["J22"].font = font(8, color=GRAY)
    helpers = {
        "J14": '=IF(B14="理系","理系","文系")',
        "J26": '=IFERROR(IF(INDEX(志望校文系,MATCH(B11,志望校リスト,0))="","",INDEX(志望校文系,MATCH(B11,志望校リスト,0))),"")',
        "K26": '=IFERROR(IF(INDEX(志望校理系,MATCH(B11,志望校リスト,0))="","",INDEX(志望校理系,MATCH(B11,志望校リスト,0))),"")',
        "J28": '=IF(ISNUMBER(B15),IF(B15>=難関閾値,3,IF(B15>=標準閾値,2,1)),IF(ISNUMBER(B26),IF(B26>=難関閾値,3,IF(B26>=標準閾値,2,1)),2))',
        "J29": '=IF(ISNUMBER(B26),IF(B26>=難関閾値,3,IF(B26>=標準閾値,2,1)),J28)',
        "J30": "=MAX(J28,MIN(J29,J28+1))",
        "J32": '=IF(B11="","志望校",B11)',
        "J33": '=IF(B5="","",B5)',
    }
    for ref, f in helpers.items():
        ws[ref] = f
        ws[ref].font = font(8, color=GRAY)
    for col in "JK":
        ws.column_dimensions[col].hidden = False
    names = {"kName": "$B$5", "kHighschool": "$B$6", "kGrade": "$B$7", "kTantou": "$B$8",
             "kSchool1": "$B$11", "kSchool2": "$B$12", "kSchool3": "$B$13",
             "kBunri": "$J$14", "kCurrent": "$B$15", "kSocial1": "$B$17", "kSocial2": "$B$18",
             "kScience": "$B$19", "kKyoteUse": "$B$20", "kAutoApply": "$B$21",
             "kKubun": "$B$24", "kSystem": "$B$25", "kTarget": "$B$26", "kPreLevel": "$B$28",
             "kPostLevel": "$B$29", "kPattern": "$B$32", "kSchool": "$J$32",
             "kAutoInputs": "$B$11:$B$20"}
    for n, ref in names.items():
        name(n, f"'生徒情報'!{ref}")
    lists = [("B7", '"高2,高3,既卒"'), ("B14", '"文系,理系"'),
             ("B17", '"日本史,世界史,地理,公共・政経,公共・倫理,数学"'),
             ("B18", '"なし,日本史,世界史,地理,公共・政経,公共・倫理"'),
             ("B19", '"化学基礎・生物基礎,物理基礎・化学基礎,物理基礎・生物基礎,物理・化学,化学・生物,物理・生物"'),
             ("B20", '"する,しない"'), ("B21", '"する,しない"'), ("B11:B13", "=志望校リスト")]
    for ref, f in lists:
        dv = DataValidation(type="list", formula1=f, allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(ref)
    dv = DataValidation(type="decimal", operator="between", formula1="20", formula2="85", allow_blank=True,
                        error="偏差値は 20〜85 の数で入力してください", showErrorMessage=True)
    ws.add_data_validation(dv)
    dv.add("B15:B16")
    # ボタンの説明
    notes = [(14, "← ボタン：管理表を、いまの判定で作り直します"),
             (15, "（チェック済み☑があるときは確認が出ます）")]
    for r, t in notes:
        ws.cell(r, 5, t).font = font(8, color=GRAY)
    ws.freeze_panes = "A4"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = 9
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 1
    ws.print_area = "A1:C32"


# ============================================================ マスタ類
def header_row(ws, headers, widths, color=NAVY):
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(1, i, h)
        c.font, c.fill = font(10, True, "FFFFFF"), fill(color)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[L(i)].width = w
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"
    ws.sheet_properties.tabColor = "BEC3CD"


def body_font(ws, r, n, sz=9):
    for i in range(1, n + 1):
        c = ws.cell(r, i)
        c.font = font(sz)
        c.alignment = Alignment(vertical="center", wrap_text=False)
        c.border = Border(bottom=THIN)


def build_masters(ws_univ, ws_mat, ws_course, ws_db, ws_ky, name):
    # 志望校マスタ
    header_row(ws_univ, ["大学名", "区分", "文系の目標偏差値", "理系の目標偏差値", "地域", "メモ（学部・方式など）"],
               [26, 9, 12, 12, 8, 40])
    for i, (u, kb, b, r, area) in enumerate(D.UNIVERSITIES, start=2):
        for j, v in enumerate((u, kb, b, r, area), start=1):
            ws_univ.cell(i, j, v)
        body_font(ws_univ, i, 6)
        for j in (3, 4):
            ws_univ.cell(i, j).number_format = "0.0"
            ws_univ.cell(i, j).alignment = Alignment(horizontal="center")
    n_end = 200
    for i in range(len(D.UNIVERSITIES) + 2, n_end + 1):
        for j in (1, 2, 3, 4):
            ws_univ.cell(i, j).fill = fill(YELLOW_IN) if i < len(D.UNIVERSITIES) + 12 else PatternFill()
    ws_univ["H1"] = ("※ 偏差値は大手予備校の公表値を参考にした『おおよその目安』です（学部・方式で異なる）。"
                     "毎年更新してください。行を追加すれば生徒情報の選択肢に出ます（200行目まで）。")
    ws_univ["H1"].font = font(9, color="C81E1E")
    ws_univ["H1"].alignment = Alignment(wrap_text=True, vertical="top")
    ws_univ.column_dimensions["H"].width = 50
    name("志望校リスト", f"'志望校マスタ'!$A$2:$A${n_end}")
    name("志望校区分", f"'志望校マスタ'!$B$2:$B${n_end}")
    name("志望校文系", f"'志望校マスタ'!$C$2:$C${n_end}")
    name("志望校理系", f"'志望校マスタ'!$D$2:$D${n_end}")
    dv = DataValidation(type="list", formula1='"国公立,私立"', allow_blank=True)
    ws_univ.add_data_validation(dv)
    dv.add(f"B2:B{n_end}")

    # 教材マスタ
    header_row(ws_mat, ["照合キー（管理表の文言に含まれる語）", "教材名（出版社）", "教科", "目標（周）", "使い方・時期", "レベル", "区分"],
               [26, 44, 9, 9, 40, 12, 18])
    for i, m in enumerate(D.MATERIALS, start=2):
        for j, v in enumerate(m, start=1):
            ws_mat.cell(i, j, v if v != "" else None)
        body_font(ws_mat, i, 7)
    ws_mat["I1"] = ("※ 管理表の【家】の帯に『照合キー』が含まれると、模試・自学シート②の教材リストに自動で載ります。"
                    "教材を差し替えるときは、カリキュラムDB の文言とこの表をセットで直してください。教材はすべて仮（要確認）。")
    ws_mat["I1"].font = font(9, color="C81E1E")
    ws_mat["I1"].alignment = Alignment(wrap_text=True, vertical="top")
    ws_mat.column_dimensions["I"].width = 50

    # 講座マスタ（一覧・参照用）
    header_row(ws_course, ["科目", "レベル", "区分", "文理", "講座コード", "講座名", "夏期", "直前", "出典・確認状態"],
               [10, 8, 8, 8, 10, 34, 10, 10, 18])
    r = 2
    for (subj, lv, kb, br), cdef in D.COURSES.items():
        vals = (subj, lv, "共通" if kb == "*" else kb, "共通" if br == "*" else br, cdef["code"] or "（要確認）",
                cdef["name"], ("N" + cdef["code"]) if cdef["summer"] and cdef["code"] else ("あり" if cdef["summer"] else "—"),
                ("F" + cdef["code"]) if cdef["chokuzen"] and cdef["code"] else ("あり" if cdef["chokuzen"] else "—"),
                cdef["status"])
        for j, v in enumerate(vals, start=1):
            ws_course.cell(r, j, v)
        body_font(ws_course, r, 9)
        if cdef["status"] != D.OK:
            ws_course.cell(r, 9).font = font(9, True, "C81E1E")
        r += 1
    r += 1
    ws_course.cell(r, 1, "共通テスト対策講座（9月スタート案）").font = font(10, True, NAVY)
    r += 1
    for k in D.KYOTE:
        vals = (k[0], "共テ", "共通", k[6], k[1].split("　")[0], k[1].split("　")[1], "", "", D.OK)
        for j, v in enumerate(vals, start=1):
            ws_course.cell(r, j, v)
        body_font(ws_course, r, 9)
        r += 1
    ws_course["K1"] = ("※ 『既存ファイル記載』＝試作ファイルにあった講座（出典：2027年度 高校生ウイングネット学習・講座ガイド P5・P22-23）。"
                       "『仮（要確認）』＝パンフレット未照合のため仮置き。講座名を直すときは カリキュラムDB の『講座』行も直してください。")
    ws_course["K1"].font = font(9, color="C81E1E")
    ws_course["K1"].alignment = Alignment(wrap_text=True, vertical="top")
    ws_course.column_dimensions["K"].width = 50

    # カリキュラムDB
    heads = ["キー（系統|科目|レベル|種別）", "系統", "科目", "レベル", "種別"] + [p[0] for p in D.PERIODS]
    header_row(ws_db, heads, [30, 9, 8, 6, 5] + [16] * NPER)
    for i, (sysn, subj, lv, kind, cells) in enumerate(D.build_db(), start=2):
        ws_db.cell(i, 1, f"{sysn}|{subj}|{lv}|{kind}")
        ws_db.cell(i, 2, sysn)
        ws_db.cell(i, 3, subj)
        ws_db.cell(i, 4, lv)
        ws_db.cell(i, 5, kind)
        for c, v in enumerate(cells):
            ws_db.cell(i, 6 + c, v or None)
        body_font(ws_db, i, 5 + NPER, 8)
        if kind == "講座":
            for c in range(1, 6 + NPER):
                ws_db.cell(i, c).border = Border(top=Side(style="thin", color=NAVY), bottom=THIN)
    for c in range(6, 6 + NPER):
        ws_db.cell(1, c).fill = fill("1F4E9E" if c - 6 < D.SPLIT else "3462B0")
    ws_db.freeze_panes = "F2"
    ws_db.auto_filter.ref = f"A1:{L(5 + NPER)}{ws_db.max_row}"
    ws_db.sheet_properties.tabColor = "E88A12"

    # 共テDB
    header_row(ws_ky, ["科目", "講座名", "9月", "10月", "11月", "12月", "対象（全/文系/理系/公共政経/公共倫理/情報）"],
               [10, 38, 9, 9, 9, 9, 30])
    for i, k in enumerate(D.KYOTE, start=2):
        for j, v in enumerate(k, start=1):
            ws_ky.cell(i, j, v or None)
        body_font(ws_ky, i, 7)


# ============================================================ 管理表
def build_plan(ws, old):
    cmap = lambda c: c if c == 1 else c + 1  # noqa: E731
    in_body = lambda r, c: FIRST <= r <= LAST and c >= 2  # noqa: E731
    copy_sheet(old, ws, cmap=cmap,
               keep_value=lambda r, c: (not in_body(r, c) and (r != 7 or c == 1 or c >= 20)
                                        and not (r <= 3 and c >= 4) and c < 26),
               keep_merge=lambda m: not (m.min_row >= FIRST and m.max_row <= LAST and m.min_col >= 2))
    # 本体（8〜79行）の書式：英語ブロック（8〜13行）を型に、科目の色だけ変える
    for s, (subj, color) in enumerate(D.SUBJECTS):
        r0 = FIRST + s * BLOCK
        for off in range(BLOCK):
            r = r0 + off
            for nc in range(2, 27):
                srcc = 2 if nc == 2 else nc - 1
                cell = ws.cell(r, nc)
                cell._style = copy(old.cell(FIRST + off, srcc)._style)
                if 2 <= nc <= 1 + NPER:      # 文字の書式は各行のB列（結合の内側は既定書式のため）
                    tpl = old.cell(FIRST + off, 2)
                    cell.font, cell.alignment = copy(tpl.font), copy(tpl.alignment)
                elif nc == COL_KY:
                    tpl = old.cell(FIRST, COL_KY - 1)
                    cell.font, cell.alignment = copy(tpl.font), copy(tpl.alignment)
                elif nc > COL_KY:
                    tpl = old.cell(FIRST, COL_KY)
                    cell.font, cell.alignment = copy(tpl.font), copy(tpl.alignment)
                if 2 <= nc <= 1 + NPER:
                    if off in (0, 2, 4):
                        cell.fill = PatternFill()
                    elif off == 1:
                        cell.fill = fill(tint(color, 0.78))
                    elif off == 3:
                        cell.fill = fill(tint(color, 0.85))
                    else:
                        cell.fill = fill("FDF3E8")
                elif nc >= COL_KY:
                    cell.fill = PatternFill()
        ws.cell(r0, 1)._style = copy(old.cell(r0, 1)._style)
        add_block_cf(ws, r0, color)
        # 進捗（S列）
        k = r0 + 1
        rng = f"$B{k}:${L(1 + NPER)}{k}"
        total = f'(COUNTIF({rng},"?*")-COUNTIF({rng},"※")+COUNTIF({rng},"*・*"))'
        done = f'(COUNTIF({rng},"☑*")+COUNTIF({rng},"☑*・*"))'
        late = (f'SUMPRODUCT((LEFT({rng},1)<>"☑")*({rng}<>"")*({rng}<>"※")*(期間終了日<基準日)'
                f'*(1+ISNUMBER(FIND("・",{rng}))))')
        ws.cell(r0, COL_PROG).value = (f'=IF({total}=0,"",{done}&"/"&{total}&"回"&IF({late}>0,CHAR(10)&"遅れ"&{late},""))')
        ws.cell(r0 + 1, COL_PROG).value = f'=IF({total}=0,"",{done}/{total})'
        for off, lab in ((3, "塾"), (5, "家")):
            rr = r0 + off
            rg = f"$B{rr}:${L(1 + NPER)}{rr}"
            ws.cell(rr, COL_PROG).value = f'=IF(COUNTIF({rg},"?")=0,"","{lab} "&COUNTIF({rg},"☑")&"/"&COUNTIF({rg},"?"))'
        S = L(COL_PROG)
        ws.conditional_formatting.add(f"{S}{r0}", FormulaRule(formula=[f'ISNUMBER(SEARCH("遅れ",${S}${r0}))'],
                                                               font=Font(color="C81E1E"), stopIfTrue=True))
        ws.conditional_formatting.add(f"{S}{r0 + 1}", DataBarRule(start_type="num", start_value=0, end_type="num",
                                                                   end_value=1, color="78A0DC"))
    # 見出し
    ws["A1"].value = "合格一貫\nカリキュラム"
    ws["E1"] = "=IF(kGrade=\"\",\"高2\",kGrade)"
    ws["E2"] = ('=IF(LEFT(kSystem,2)="国立","国公立大","私立大")&"（"&kBunri&"）"&CHAR(10)&"前期 "&kPreLevel&" → 後期 "&kPostLevel')
    ws["I1"] = ('="第1志望："&IF(kSchool1="","（未入力）",kSchool1)&IF(kSchool2="",""," ／ 第2志望："&kSchool2)'
                '&IF(kSchool3="",""," ／ 第3志望："&kSchool3)&CHAR(10)&"目標偏差値："&IF(kTarget="","—",TEXT(kTarget,"0.0"))'
                '&"　現在："&IF(kCurrent="","—",TEXT(kCurrent,"0.0"))')
    ws["U1"], ws["U2"], ws["U3"] = "氏名", "高校", "担当"
    ws["V1"] = '=IF(kName="","",kName)'
    ws["V2"] = '=IF(kHighschool="","",kHighschool)'
    ws["V3"] = '=IF(kTantou="","",kTantou)'
    for ref in ("U1", "U2", "U3"):
        ws[ref]._style = copy(old["T1"]._style)
    for ref in ("V1", "V2", "V3"):
        ws[ref]._style = copy(old["U1"]._style)
    ws["V1"].font = font(14, True, INK)
    # 4行目：学年の帯
    ws.row_dimensions[4].height = 13
    ws.merge_cells("B4:D4")
    ws.merge_cells(f"E4:{L(1 + NPER)}4")
    for ref, txt, col in (("B4", "◀ 高2 ▶", "3462B0"), ("E4", "◀ 高3 ▶", NAVY)):
        ws[ref] = txt
        ws[ref].font, ws[ref].fill = font(8, True, "FFFFFF"), fill(col)
        ws[ref].alignment = Alignment(horizontal="center", vertical="center")
    # 6行目：学習目的
    ws["B6"] = "入試基礎力を強化する（高2の1月〜）"
    # 7行目：期間
    for i, (p, _, _) in enumerate(D.PERIODS):
        c = ws.cell(7, 2 + i, p)
        src = old["D7"] if p in D.HIGHLIGHT_PERIODS else old["C7"]
        c._style = copy(src._style)
    # 共通テスト側
    ws["U5"] = "大学入学共通テスト対策（9月スタート）"
    dv = DataValidation(type="list", formula1='"✓"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"U{FIRST}:U{LAST}")
    # 説明・フッター
    ws[f"A{LAST + 1}"] = (
        "【使い方】各科目6行：講座の帯／回（ダブルクリックで☑）／【塾】授業の教材／☐毎月の進み（ダブルクリックで☑）／【家】家庭学習の教材／☐毎月の進み。"
        "共通テスト対策は9月スタート（右欄の回もダブルクリックで☑）。終了日（設定シート）を過ぎても未完了は赤。"
        "内容は『生徒情報』の志望校・偏差値から自動で組み立て（前期＝現在の偏差値、後期＝志望校の目標偏差値でレベルを選択）。"
        "期間は高2の1月〜直前Ⅱ期（翌2月）。※＝推奨受講期間。講座の出典：2027年度 高校生ウイングネット学習・講座ガイド P5・P22-23（一部は仮置き：講座マスタ参照）。教材は仮（要確認）。")
    ws[f"V{LAST + 2}"] = '="作成日："&TEXT(基準日,"yyyy年m月d日")'
    for r, t in ((4, "↑ 押すと成績表（A4縦）の印刷プレビューが開きます"),
                 (9, "↑ 設定シートの受験科目（○）に合わせて科目を表示/非表示"),
                 (16, "↑ 生徒情報の内容で管理表を作り直します")):
        c = ws.cell(r, 28, t)
        c.font = font(8, color=GRAY)
    ws.column_dimensions["AB"].width = 8.6
    for col in ("AC", "AD", "AE"):
        ws.column_dimensions[col].width = 8.6
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = 8
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 1
    ws.sheet_properties.tabColor = NAVY


def add_block_cf(ws, r0, color):
    end = L(1 + NPER)
    white_b = Font(b=True, color="FFFFFF")
    band = DifferentialStyle(font=white_b, fill=PatternFill(bgColor=color, fill_type="solid"))
    from openpyxl.formatting.rule import Rule
    ws.conditional_formatting.add(f"B{r0}:{end}{r0}", Rule(type="expression", dxf=band, formula=[f'B{r0}<>""']))
    # 回
    k = r0 + 1
    ws.conditional_formatting.add(f"B{k}:{end}{k}", Rule(type="expression", stopIfTrue=True, dxf=DifferentialStyle(
        font=Font(b=True, color="FFFFFF"), fill=PatternFill(bgColor=color, fill_type="solid")), formula=[f'LEFT(B{k},1)="☑"']))
    ws.conditional_formatting.add(f"B{k}:{end}{k}", Rule(type="expression", stopIfTrue=True, dxf=DifferentialStyle(
        font=Font(b=True, color="C81E1E"), fill=PatternFill(bgColor="FFCDCD", fill_type="solid")),
        formula=[f'AND(B{k}<>"",B{k}<>"※",LEFT(B{k},1)<>"☑",INDEX(期間終了日,COLUMN()-1)<基準日)']))
    ws.conditional_formatting.add(f"B{k}:{end}{k}", Rule(type="expression", stopIfTrue=True, dxf=DifferentialStyle(
        font=Font(color="BEC3CD")), formula=[f'B{k}="※"']))
    # 教材の帯
    for off, bg in ((2, tint(color, 0.55)), (4, "FAE5CB")):
        rr = r0 + off
        ws.conditional_formatting.add(f"B{rr}:{end}{rr}", Rule(type="expression", stopIfTrue=True, dxf=DifferentialStyle(
            font=Font(color="9AA0AE"), fill=PatternFill(bgColor="F4F5F8", fill_type="solid")), formula=[f'LEFT(B{rr},1)="（"']))
        ws.conditional_formatting.add(f"B{rr}:{end}{rr}", Rule(type="expression", dxf=DifferentialStyle(
            fill=PatternFill(bgColor=bg, fill_type="solid")), formula=[f'B{rr}<>""']))
    # ☐ 行
    for off in (3, 5):
        rr = r0 + off
        ws.conditional_formatting.add(f"B{rr}:{end}{rr}", Rule(type="expression", stopIfTrue=True, dxf=DifferentialStyle(
            font=Font(b=True, color="1E8A4A")), formula=[f'B{rr}="☑"']))
        ws.conditional_formatting.add(f"B{rr}:{end}{rr}", Rule(type="expression", stopIfTrue=True, dxf=DifferentialStyle(
            font=Font(b=True, color="C81E1E"), fill=PatternFill(bgColor="FFCDCD", fill_type="solid")),
            formula=[f'AND(B{rr}="☐",INDEX(期間終了日,COLUMN()-1)<基準日)']))
    # 共通テスト
    rng = f"V{r0}:Z{r0 + BLOCK - 1}"
    ws.conditional_formatting.add(rng, Rule(type="expression", stopIfTrue=True, dxf=DifferentialStyle(
        font=Font(b=True, color="FFFFFF"), fill=PatternFill(bgColor=color, fill_type="solid")), formula=[f'LEFT(V{r0},1)="☑"']))
    ws.conditional_formatting.add(rng, Rule(type="expression", dxf=DifferentialStyle(
        fill=PatternFill(bgColor=tint(color, 0.75), fill_type="solid")), formula=[f'V{r0}<>""']))


# ============================================================ 模試・自学
def build_mock(ws, old):
    rmap = lambda r: r if r <= 27 else r + 2  # noqa: E731
    copy_sheet(old, ws, rmap=rmap, keep_value=lambda r, c: not (20 <= r <= 27) and c != 10 and not (r in (5, 14, 15)))
    for r in (28, 29):     # 教材行を2行追加
        ws.merge_cells(f"F{r}:H{r}")
    for r, srow in ((27, 26), (28, 26), (29, 27)):
        for c in range(1, 11):
            ws.cell(r, c)._style = copy(old.cell(srow, c)._style)
        ws.row_dimensions[r].height = old.row_dimensions[srow].height
    for r in range(MAT_FIRST, MAT_FIRST + MAT_N):
        ws.cell(r, 5).value = f'=IF(OR(C{r}="",D{r}=""),"",MIN(1,D{r}/C{r}))'
        for c in (1, 2, 3, 4, 6):
            ws.cell(r, c).value = None
    ws["A5"] = "目標（志望校）"
    ws["H5"] = '=IF(kTarget="","",kTarget)'
    ws["J3"], ws["J4"], ws["J6"] = "集計", "模試到達度", "自学進捗"
    ws["J5"] = '=IF(OR(H14="",H5="",H5<=偏差値基準),"",MAX(0,MIN(1,(H14-偏差値基準)/(H5-偏差値基準))))'
    ws["J7"] = f"=IFERROR(AVERAGE(E{MAT_FIRST}:E{MAT_FIRST + MAT_N - 1}),\"\")"
    ws["A14"], ws["A15"] = "最新", "目標との差"
    ws["B14"] = '=IFERROR(LOOKUP(2,1/($A$6:$A$13<>""),$B$6:$B$13),"")'
    for col in "CDEFGH":
        ws[f"{col}14"] = f'=IFERROR(LOOKUP(2,1/({col}$6:{col}$13<>""),{col}$6:{col}$13),"")'
        ws[f"{col}15"] = f'=IF(OR({col}14="",{col}5=""),"",{col}14-{col}5)'
    ws["A18"] = "② 家庭学習推奨教材（全教科）　※ おすすめカリキュラムの反映で自動入力。『現在（周）』を毎回入力"
    ws["A2"] = "※ 目標（志望校）の総合は生徒情報（志望校マスタ）から自動で入ります。教科別の目標は手入力"
    ws["A2"].font = font(8, color=GRAY)
    ws.print_area = "A1:J42"
    ws.page_setup.orientation, ws.page_setup.paperSize = "portrait", 9
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 1
    copy_cf(old, ws, rmap=rmap)


# ============================================================ 成績表
def build_report(ws, old):
    rmap = lambda r: r if r <= 47 else r + 2  # noqa: E731
    copy_sheet(old, ws, rmap=rmap)
    for r in (48, 49):
        for a, b in (("A", "B"), ("C", "F"), ("G", "H"), ("I", "L")):
            ws.merge_cells(f"{a}{r}:{b}{r}")
    for r, srow in ((47, 46), (48, 46), (49, 47)):
        for c in range(1, 17):
            ws.cell(r, c)._style = copy(old.cell(srow, c)._style)
        ws.row_dimensions[r].height = old.row_dimensions[srow].height
    plan = PLAN
    ws["C3"] = '=IF(kName="","（未入力）",kName&"　さん")'
    ws["I3"] = '=IF(kHighschool="","",kHighschool&"　")&kGrade'
    ws["C4"] = ('=IF(LEFT(kSystem,2)="国立","国公立大","私立大")&"（"&kBunri&"）　前期 "&kPreLevel&" → 後期 "&kPostLevel')
    ws["I4"] = '=IF(kTantou="","",kTantou)'
    ws["C5"] = f"={plan}!I1"
    ws["A8"] = '=IFERROR(AVERAGE(F9,I9,L9),"")'
    ws["E8"] = '=IF(SUM(O12:O23)=0,"—",SUM(N12:N23)&" / "&SUM(O12:O23)&" 回")'
    ws["H8"] = '=IF(模試・自学!H14="","—","最新 "&TEXT(模試・自学!H14,"0.0")&" / 目標 "&TEXT(模試・自学!H5,"0.0"))'
    ws["F9"] = '=IF(SUM(O12:O23)=0,"",SUM(N12:N23)/SUM(O12:O23))'
    ws["I9"] = "=模試・自学!J5"
    ws["L9"] = "=模試・自学!J7"
    end = L(1 + NPER)
    for i in range(NSUB):
        r = 12 + i
        k = FIRST + i * BLOCK + 1
        rng = f"{plan}!$B${k}:${end}${k}"
        on = f'(設定!$B${SET_FIRST + i}="○")'
        ws[f"A{r}"] = f"={plan}!A{FIRST + i * BLOCK}"
        ws[f"C{r}"] = f'=IF(O{r}=0,"",N{r}&" / "&O{r}&" 回")'
        ws[f"E{r}"] = f'=IF(O{r}=0,"",N{r}/O{r})'
        ws[f"J{r}"] = f'=IF(O{r}=0,"",IF(P{r}=0,"予定どおり","遅れ "&P{r}&" 回"))'
        ws[f"N{r}"] = f'=((COUNTIF({rng},"☑*")+COUNTIF({rng},"☑*・*")))*{on}'
        ws[f"O{r}"] = f'=((COUNTIF({rng},"?*")-COUNTIF({rng},"※")+COUNTIF({rng},"*・*")))*{on}'
        ws[f"P{r}"] = (f'=(SUMPRODUCT((LEFT({rng},1)<>"☑")*({rng}<>"")*({rng}<>"※")*(期間終了日<基準日)'
                       f'*(1+ISNUMBER(FIND("・",{rng})))))*{on}')
    for i in range(8):
        r, m = 27 + i, 6 + i
        for col, mc in zip("ADFGHIJK", "ABCDEFGH"):
            ws[f"{col}{r}"] = f'=IF(模試・自学!{mc}{m}="","",模試・自学!{mc}{m})'
    for col, mc in zip("FGHIJK", "CDEFGH"):
        ws[f"{col}35"] = f'=IF(模試・自学!{mc}5="","",模試・自学!{mc}5)'
        ws[f"{col}36"] = f"=模試・自学!{mc}15"
    for i in range(MAT_N):
        r, m = 40 + i, MAT_FIRST + i
        ws[f"A{r}"] = f'=IF(模試・自学!A{m}="","",模試・自学!B{m})'
        ws[f"C{r}"] = f'=IF(模試・自学!A{m}="","",模試・自学!A{m})'
        ws[f"G{r}"] = (f'=IF(模試・自学!A{m}="","",IF(模試・自学!D{m}="","-",模試・自学!D{m})&" / "'
                       f'&IF(模試・自学!C{m}="","-",模試・自学!C{m})&" 周")')
        ws[f"I{r}"] = f"=模試・自学!E{m}"
    ws["A52"] = '=IF(模試・自学!A33="","",模試・自学!A33)'
    ws["A58"] = '=IF(模試・自学!A40="","",模試・自学!A40)'

    copy_cf(old, ws, rmap=rmap)
    # 教材の進捗バーを2行分のばす
    ws.conditional_formatting.add("I48:I49", DataBarRule(start_type="num", start_value=0, end_type="num",
                                                          end_value=1, color="78A0DC"))


# ============================================================ 使い方
def build_howto(ws):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 110
    ws.sheet_properties.tabColor = "5CAC4A"
    lines = [
        ("合格一貫カリキュラム　マスターテンプレート（生徒1人につき1ファイル）", "title"),
        ("このファイルを生徒ごとにコピーして使います（例：2027_山田太郎.xlsm）。マクロを有効にして開いてください。", "note"),
        ("1. 生徒情報シート", "h"),
        ("黄色のセルに 氏名・学年・第1〜第3志望校・文理・現在の偏差値・地歴公民・理科 を入力。", ""),
        ("第1志望校（志望校マスタ）から『国公立／私立』と目標偏差値が決まり、系統（国立文系・国立理系・私大文系・私大理系）を自動判定。", ""),
        ("前期（高2の1月〜夏期）のレベルは現在の偏差値、後期（9月〜直前）のレベルは志望校の目標偏差値で『基礎／標準／難関』から選ばれます。", ""),
        ("入力を変えると、おすすめカリキュラムが管理表に自動で反映されます（☑ の記録があるときは確認が出ます）。", ""),
        ("2. 合格一貫カリキュラム（管理表・A3横で印刷）", "h"),
        ("各科目6行：講座の帯／回／【塾】授業の教材／☐（毎月の進み）／【家】家庭学習の教材／☐（毎月の進み）。期間は高2の1月〜直前Ⅱ期（翌2月）の17区分。", ""),
        ("受講した回・進んだ月は、セルをダブルクリック（タブレットはダブルタップ）で ☑。もう一度で外れます。右側は共通テスト対策（9月〜12月）。", ""),
        ("設定シートの終了日を過ぎても ☑ がない回・☐ は赤くなります。S列に科目ごとの進捗（回数・遅れ・％）。", ""),
        ("講座や教材を個別に変えたいときは、管理表のセルを直接書き換えてかまいません（再反映すると元に戻ります）。", ""),
        ("3. 模試・自学シート", "h"),
        ("模試の偏差値を古い順に入力。② の家庭学習推奨教材は反映時に自動で並ぶので、『現在（周）』だけ毎回入力。③講師コメント・④次回の目標も入力。", ""),
        ("4. 成績表（A4縦）", "h"),
        ("『成績表を印刷』ボタンで、使っていない行を隠して印刷プレビューを開きます。総合達成度＝講座の受講・模試・家庭学習教材の平均。", ""),
        ("5. 設定シート", "h"),
        ("年度（生徒の高3の年度）・各期間の終了日・レベル判定の基準偏差値（難関60／標準50）・受験科目の○を管理します。", ""),
        ("6. マスタ（志望校・教材・講座・カリキュラムDB・共テDB）", "h"),
        ("カリキュラムDB：系統×科目×レベル×種別（講座/回/塾/家）の1行が17期間分。『←』は左のセルと結合（帯をのばす）。{志望校} は第1志望校名に置き換わります。", ""),
        ("前期（1月〜夏期Ⅱ期）と後期（9月〜直前Ⅱ期）は別々のレベルから組み合わせるので、帯は夏期Ⅱ期と9月の間をまたがないようにしてください。", ""),
        ("講座名・教材・偏差値には仮置きを含みます（講座マスタ・教材マスタの確認状態を参照）。パンフレットと照合して直してください。", "note"),
    ]
    r = 1
    for text, kind in lines:
        c = ws.cell(r, 2, text)
        if kind == "title":
            c.font = font(15, True, "FFFFFF")
            c.fill = fill(NAVY)
            ws.row_dimensions[r].height = 30
        elif kind == "h":
            c.font = font(11, True, NAVY)
            c.fill = fill(PALE)
            r += 0
            ws.row_dimensions[r].height = 20
        elif kind == "note":
            c.font = font(9.5, True, "C81E1E")
        else:
            c.font = font(10)
        c.alignment = Alignment(vertical="center", wrap_text=True, indent=1)
        r += 1


if __name__ == "__main__":
    src = sys.argv[1]
    out = sys.argv[2]
    build(src, out)
    print("written", out)
