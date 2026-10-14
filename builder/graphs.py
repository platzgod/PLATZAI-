"""成績グラフ（学習達成度レポートの2ページ目・A4縦）。

  ① 達成度：総合・講座の受講・模試・家庭学習教材 のドーナツ4つ（中の％はセルに大きく表示）
  ② 月別の受講回数：期間ごとの予定回数と受講済み回数（集合縦棒）
  ③ 模試の偏差値の推移：総合・英語・数学・国語（折れ線。未入力の回は線を切る）

グラフの元データは、印刷範囲の外（N列〜）に数式で置く。マクロなしでも数式だけで更新される。
色はデータ可視化の検証済みパレット（青・オレンジ・アクア・黄）を使用。
"""
from openpyxl.chart import BarChart, DoughnutChart, LineChart, Reference, Series
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.chart.marker import Marker
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import CharacterProperties, Font as DFont, Paragraph, ParagraphProperties
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

import data as D

SHEET = "成績グラフ"
FONT = "游ゴシック"
NAVY = "17358F"
INK = "1D2433"
GRAY = "6B7488"
GRID = "E3E6EC"
TRACK = "E6E9EF"          # ドーナツの残り
BLUE, ORANGE, AQUA, YELLOW = "2A78D6", "EB6834", "1BAF7A", "EDA100"
PLAN_GRAY = "B7C2D3"      # 予定（背景的な棒）
PLAN = "合格一貫カリキュラム"
HC = 14                   # 補助データの先頭列（N）


def _font(sz, b=False, color=INK):
    return Font(name=FONT, sz=sz, b=b, color=color)


def _txt(sz=800, color=GRAY, b=False):
    cp = CharacterProperties(sz=sz, b=b, solidFill=color, latin=DFont(typeface=FONT), ea=DFont(typeface=FONT))
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])


def _axis_style(ax, grid=False):
    ax.delete = False
    ax.txPr = _txt()
    ax.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill="C9CED8", w=6350))
    if grid:
        ax.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill=GRID, w=6350))
    else:
        ax.majorGridlines = None


def place(ws, chart, c1, r1, c2, r2):
    """セル範囲（0始まり：c1,r1 から c2,r2 まで）にぴったり置く。枠線なし"""
    chart.anchor = TwoCellAnchor(_from=AnchorMarker(col=c1, row=r1), to=AnchorMarker(col=c2, row=r2))
    chart.graphical_properties = GraphicalProperties(ln=LineProperties(noFill=True))
    ws.add_chart(chart)


def section(ws, r, text):
    ws.merge_cells(f"A{r}:L{r}")
    c = ws.cell(r, 1, text)
    c.font, c.fill = _font(10.5, True, "FFFFFF"), PatternFill("solid", fgColor="1F4E9E")
    c.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[r].height = 20


def build(ws):
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = "2A78D6"
    for i in range(1, 13):
        ws.column_dimensions[L(i)].width = 8.4
    ws.column_dimensions["M"].width = 3
    for i in range(HC, HC + 4):
        ws.column_dimensions[L(i)].width = 11

    # 見出し
    ws.merge_cells("A1:H1")
    ws["A1"] = "学習達成度グラフ"
    ws["A1"].font = _font(18, True, NAVY)
    ws["A1"].alignment = Alignment(vertical="center")
    ws.merge_cells("I1:L1")
    ws["I1"] = "個人指導塾 PLATZ"
    ws["I1"].font = _font(11, True, NAVY)
    ws["I1"].alignment = Alignment(horizontal="right", vertical="center")
    ws.row_dimensions[1].height = 34
    ws.merge_cells("A2:H2")
    ws["A2"] = '=IF(kName="","",kName&"　さん")&"　／　第1志望："&IF(kSchool1="","（未入力）",kSchool1)'
    ws["A2"].font = _font(10.5, True, INK)
    ws.merge_cells("I2:L2")
    ws["I2"] = '="作成日："&TEXT(基準日,"yyyy年m月d日")'
    ws["I2"].font = _font(9, False, GRAY)
    ws["I2"].alignment = Alignment(horizontal="right")
    for c in range(1, 13):
        ws.cell(2, c).border = Border(bottom=Side(style="medium", color=NAVY))

    # ---------------------------------------------------------- 補助データ（印刷範囲外）
    hdr = ws.cell(3, HC, "グラフ用データ（自動計算・編集不要）")
    hdr.font = _font(8, True, GRAY)
    # ① 達成度
    gauges = [("総合達成度", "成績表!A8"), ("講座の受講", "成績表!F9"), ("模試（偏差値）", "成績表!I9"), ("家庭学習教材", "成績表!L9")]
    for i, (lab, ref) in enumerate(gauges):
        r = 5 + i
        ws.cell(r, HC, lab)
        ws.cell(r, HC + 1, f'=IF(ISNUMBER({ref}),MAX(0,MIN(1,{ref})),0)').number_format = "0%"
        ws.cell(r, HC + 2, f"=1-{L(HC + 1)}{r}").number_format = "0%"
        ws.cell(r, HC + 3, f'=IF(ISNUMBER({ref}),"","未入力")')
    # ② 月別
    m0 = 12
    ws.cell(m0 - 1, HC, "期間")
    ws.cell(m0 - 1, HC + 1, "予定（回）")
    ws.cell(m0 - 1, HC + 2, "受講済み（回）")
    for p, (name, _, _) in enumerate(D.PERIODS):
        r = m0 + p
        col = L(2 + p)
        rng = f"{PLAN}!{col}$8:{col}$79"
        kai = f"(MOD(ROW({rng})-9,6)=0)"
        dbl = f'(1+ISNUMBER(FIND("・",{rng})))'
        ws.cell(r, HC, name)
        ws.cell(r, HC + 1, f'=SUMPRODUCT({kai}*({rng}<>"")*({rng}<>"※")*{dbl})')
        ws.cell(r, HC + 2, f'=SUMPRODUCT({kai}*(LEFT({rng},1)="☑")*{dbl})')
    for r in range(4, m0 + D.NPER):
        for c in range(HC, HC + 4):
            ws.cell(r, c).font = _font(8, False, GRAY)

    # ---------------------------------------------------------- ① 達成度（ドーナツ）
    section(ws, 4, "① 志望校までの達成度")
    for i, (lab, _) in enumerate(gauges):
        r = 5 + i
        ch = DoughnutChart(holeSize=62)
        ch.add_data(Reference(ws, min_col=HC + 1, max_col=HC + 2, min_row=r, max_row=r), from_rows=True, titles_from_data=False)
        s = ch.series[0]
        from openpyxl.chart.series import DataPoint
        for idx, colr in ((0, BLUE if i == 0 else "4A90DA"), (1, TRACK)):
            pt = DataPoint(idx=idx)
            pt.graphicalProperties = GraphicalProperties(solidFill=colr, ln=LineProperties(solidFill="FFFFFF", w=19050))
            s.dPt.append(pt)
        ch.legend = None
        ch.title = None
        ch.layout = Layout(manualLayout=ManualLayout(x=0.12, y=0.04, w=0.76, h=0.92))
        place(ws, ch, 3 * i, 4, 3 * i + 3, 12)
        # ％と見出し（グラフの下のセル）
        c1 = L(1 + 3 * i)
        c3 = L(3 + 3 * i)
        ws.merge_cells(f"{c1}13:{c3}13")
        big = ws[f"{c1}13"]
        big.value = f'=TEXT({L(HC + 1)}{r},"0%")&IF({L(HC + 3)}{r}="",""," ")&{L(HC + 3)}{r}'
        big.font = _font(18 if i == 0 else 15, True, NAVY if i == 0 else INK)
        big.alignment = Alignment(horizontal="center", vertical="center")
        ws.merge_cells(f"{c1}14:{c3}14")
        lc = ws[f"{c1}14"]
        lc.value = lab
        lc.font = _font(9.5, True, GRAY)
        lc.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[13].height = 26
    ws.merge_cells("A15:L15")
    ws["A15"] = "※ 総合達成度＝講座の受講・模試・家庭学習教材の平均（未入力の項目は除く）。模試は偏差値40を0%、志望校の目標偏差値を100%として計算。"
    ws["A15"].font = _font(7.5, False, GRAY)

    # ---------------------------------------------------------- ② 月別の受講回数
    section(ws, 17, "② 期間ごとの受講回数（予定と受講済み）")
    bc = BarChart()
    bc.type = "col"
    bc.grouping = "clustered"
    bc.gapWidth = 60
    bc.overlap = -5
    data = Reference(ws, min_col=HC + 1, max_col=HC + 2, min_row=m0 - 1, max_row=m0 + D.NPER - 1)
    cats = Reference(ws, min_col=HC, min_row=m0, max_row=m0 + D.NPER - 1)
    bc.add_data(data, titles_from_data=True)
    bc.set_categories(cats)
    for s, colr in zip(bc.series, (PLAN_GRAY, BLUE)):
        s.graphicalProperties = GraphicalProperties(solidFill=colr, ln=LineProperties(noFill=True))
    _axis_style(bc.x_axis)
    _axis_style(bc.y_axis, grid=True)
    bc.y_axis.numFmt = "0"
    bc.y_axis.scaling.min = 0
    bc.legend.position = "t"
    bc.legend.txPr = _txt(850, INK)
    place(ws, bc, 0, 17, 12, 31)

    # ---------------------------------------------------------- ③ 模試の偏差値
    section(ws, 33, "③ 模試の偏差値の推移")
    ws.merge_cells("A34:L34")
    ws["A34"] = '="目標（総合）："&IF(模試・自学!H5="","—",TEXT(模試・自学!H5,"0.0"))&"　／　最新（総合）："&IF(模試・自学!H14="","—",TEXT(模試・自学!H14,"0.0"))'
    ws["A34"].font = _font(9.5, True, INK)
    lc = LineChart()
    for col, colr, title in (("H", BLUE, "総合"), ("C", ORANGE, "英語"), ("D", AQUA, "数学"), ("E", YELLOW, "国語")):
        ci = ord(col) - 64
        s = Series(Reference(_mock_sheet(ws), min_col=ci, min_row=6, max_row=13), title=title)
        lc.series.append(s)
        s.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill=colr, w=25400 if col == "H" else 19050))
        s.marker = Marker(symbol="circle", size=7)
        s.marker.graphicalProperties = GraphicalProperties(solidFill=colr, ln=LineProperties(solidFill="FFFFFF", w=9525))
        s.smooth = False
    lc.set_categories(Reference(_mock_sheet(ws), min_col=1, min_row=6, max_row=13))
    lc.display_blanks = "gap"
    _axis_style(lc.x_axis)
    _axis_style(lc.y_axis, grid=True)
    lc.y_axis.numFmt = "0"
    lc.legend.position = "t"
    lc.legend.txPr = _txt(850, INK)
    place(ws, lc, 0, 34, 12, 50)

    ws.merge_cells("A52:L52")
    ws["A52"] = "高校・大学受験専門 個人指導塾PLATZ ／ 〒373-0817 群馬県太田市飯塚町1417-1 ／ TEL 0276-55-4230 ／ 定休日：日・水"
    ws["A52"].font = _font(8, False, GRAY)
    ws["A52"].alignment = Alignment(horizontal="center")
    ws.print_area = "A1:L52"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = 9
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 1
    ws.print_options.horizontalCentered = True


def _mock_sheet(ws):
    return ws.parent["模試・自学"]
