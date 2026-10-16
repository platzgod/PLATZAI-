"""生成した .xlsm を LibreOffice で開き、VBA マクロを実行して Python 版の反映結果と突き合わせる。

    python3 builder/test_vba.py dist/xxx.xlsm
"""
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import openpyxl  # noqa: E402

import build as B  # noqa: E402
import data as D  # noqa: E402
import handout as H  # noqa: E402
from lo_harness import Office  # noqa: E402

CASES = [
    # checks=None → 「標準パターンでチェック」ボタン（SetStandardChecks）で入れる
    dict(name="国公立大（理系）・標準チェック", school1="群馬大学", bunri="理系", checks=None),
    dict(name="国公立大（理系）・下段", school1="群馬大学", bunri="理系", tier="下段", checks=None),
    dict(name="難関私立大（文系）・世界史", school1="明治大学", bunri="文系",
         checks=dict(ko=["英語", "現代文", "古文", "世界史"], ky=[])),
    dict(name="難関私立大（理工系）・共テ利用", school1="東京理科大学", bunri="理系",
         checks=dict(ko=["英語", "数学", "化学", "生物"], ky=["英語", "数学", "化学", "生物"])),
    dict(name="難関国公立大（文系）・地理・倫理", school1="千葉大学", bunri="文系",
         checks=dict(ko=["英語", "数学", "現代文", "古文", "漢文", "地理"],
                     ky=["英語", "数学", "現代文", "古文", "漢文", "地理", "物理基礎", "化学基礎", "公共・倫理", "情報Ⅰ"])),
    dict(name="手動指定・最難関私立大（文系）", school1="明治大学", bunri="文系", pattern_manual="最難関私立大（文系）",
         checks=dict(ko=["英語", "数学", "現代文", "古文"], ky=[])),
    dict(name="最難関国公立大（文系）・標準チェック", school1="東京大学", bunri="文系", checks=None),
    dict(name="有名私立大（薬学部）・下段", school1="明治薬科大学", bunri="理系", tier="下段",
         checks=dict(ko=["英語", "数学", "化学"], ky=[])),
    dict(name="全項目チェック（DBにない組み合わせを含む）", school1="群馬大学", bunri="理系",
         checks=dict(ko=list(D.ITEM_NAMES), ky=list(D.ITEM_NAMES))),
    dict(name="チェックなし・志望校なし", school1="", bunri="文系", checks=dict(ko=[], ky=[])),
]

def case_checks(case, system):
    if case["checks"] is None:
        return B.standard_checks(system)
    ko, ky = case["checks"]["ko"], case["checks"]["ky"]
    return {it: (it in ko, it in ky) for it in D.ITEM_NAMES}


def expected(case):
    univ = B.univ_table()
    inp = dict(case)
    j = B.judge(inp, univ)
    checks = case_checks(case, j["system"])
    subjects = B.subjects_from_checks(checks)
    wb = openpyxl.Workbook()
    ws = wb.active
    db_index = {(r[0], r[1], r[2], r[3]): r[4] for r in D.build_db()}
    B.reflect(ws, j, checks, subjects, db_index, inp["school1"])
    cells = {}
    for r in range(B.FIRST, B.LAST + 1):
        for c in list(range(2, 2 + B.NPER)) + [B.COL_KY] + list(B.KY_COLS):
            v = ws.cell(r, c).value
            if v not in (None, ""):
                cells[(r, c)] = str(v)
    merges = {str(m) for m in ws.merged_cells.ranges}
    mats = [m[1] for m in B.materials_for(ws, subjects)]
    ho = wb.create_sheet("ho")
    last = H.write_rows(ho, H.handout_rows(ws, subjects, B.FIRST, B.BLOCK, B.COL_KY), H.summary_text(j, checks),
                        dict(D.SUBJECTS))
    hcells = {(r, c): str(ho.cell(r, c).value) for r in range(H.H_FIRST, last + 4) for c in range(1, H.LASTC + 1)
              if ho.cell(r, c).value not in (None, "") and not str(ho.cell(r, c).value).startswith("=")}
    hcells[(3, 1)] = ho["A3"].value
    hmerges = {str(m) for m in ho.merged_cells.ranges if m.min_row >= H.H_FIRST}
    return j, checks, subjects, cells, merges, mats, (hcells, hmerges, last)


def main(path):
    fails = 0
    with Office() as off:
        doc = off.open(path)
        sh = off.sheet(doc, B.INPUT)
        plan = off.sheet(doc, B.PLAN)
        mock = off.sheet(doc, "模試・自学")
        sett = off.sheet(doc, "設定")
        nm = lambda n: doc.NamedRanges.getByName(n).getReferredCells()  # noqa: E731
        for case in CASES:
            off.set(doc, B.INPUT, "B18", "しない")  # 自動反映を切ってから入力（イベントの影響を避ける）
            off.set(doc, B.INPUT, "B11", case["school1"])
            off.set(doc, B.INPUT, "B14", case["bunri"])
            off.set(doc, B.INPUT, "B16", case.get("pattern_manual") or "")
            off.set(doc, B.INPUT, "B17", case.get("tier") or "上段")
            doc.calculateAll()
            j, checks, subjects, cells, merges, mats, hexp = expected(case)
            if case["checks"] is None:
                off.run(doc, "modCurriculum.SetStandardChecks")
            else:
                chk = nm("kChecks")
                for i, it in enumerate(D.ITEM_NAMES):
                    chk.getCellByPosition(0, i).String = "✓" if checks[it][0] else ""
                    chk.getCellByPosition(1, i).String = "✓" if checks[it][1] else ""
            doc.calculateAll()
            off.run(doc, "modCurriculum.ApplyRecommendedSilent")
            le = off.run(doc, "modCurriculum.LastError")[0]
            if le:
                print("    VBAエラー:", le)
            doc.calculateAll()
            got_sys = nm("kSystem").getCellByPosition(0, 0).String
            got_lv = nm("kPatternKey").getCellByPosition(0, 0).String
            errs = []
            if (got_sys, got_lv) != (j["system"], j["pattern"]):
                errs.append(f"判定 {got_sys},{got_lv} != {j['system']},{j['pattern']}")
            chk = nm("kChecks")
            gchk = {it: (chk.getCellByPosition(0, i).String == "✓", chk.getCellByPosition(1, i).String == "✓")
                    for i, it in enumerate(D.ITEM_NAMES)}
            if gchk != checks:
                errs.append("チェックちがい: " + str([it for it in D.ITEM_NAMES if gchk[it] != checks[it]]))
            for i, subj in enumerate(D.SUBJECT_NAMES):
                g = (sett.getCellByPosition(1, B.SET_FIRST + i - 1).String, sett.getCellByPosition(2, B.SET_FIRST + i - 1).String)
                if g != subjects[subj]:
                    errs.append(f"科目 {subj}: {g} != {subjects[subj]}")
            got = {}
            for r in range(B.FIRST, B.LAST + 1):
                for c in list(range(2, 2 + B.NPER)) + [B.COL_KY] + list(B.KY_COLS):
                    v = plan.getCellByPosition(c - 1, r - 1).String
                    if v:
                        got[(r, c)] = v
            if got != cells:
                diff = [(k, got.get(k), cells.get(k)) for k in sorted(set(got) | set(cells)) if got.get(k) != cells.get(k)]
                errs.append(f"管理表のセルが {len(diff)} 個ちがう: {diff[:5]}")
            # 結合
            gm = set()
            for r in range(B.FIRST, B.LAST + 1):
                c = 2
                while c <= 1 + B.NPER:
                    cell = plan.getCellByPosition(c - 1, r - 1)
                    cur = plan.createCursorByRange(cell)
                    cur.collapseToMergedArea()
                    a = cur.getRangeAddress()
                    if a.EndColumn > a.StartColumn:
                        gm.add(f"{openpyxl.utils.get_column_letter(a.StartColumn + 1)}{r}:{openpyxl.utils.get_column_letter(a.EndColumn + 1)}{r}")
                        c = a.EndColumn + 2
                    else:
                        c += 1
            if gm != merges:
                errs.append(f"結合ちがい: VBAのみ {sorted(gm - merges)[:4]} / Pythonのみ {sorted(merges - gm)[:4]}")
            gmats = [mock.getCellByPosition(0, B.MAT_FIRST + i - 1).String for i in range(B.MAT_N)]
            gmats = [x for x in gmats if x]
            if gmats != mats:
                errs.append(f"教材 {gmats} != {mats}")
            # 配布用カリキュラム
            hcells, hmerges, hlast = hexp
            ho = off.sheet(doc, H.SHEET)
            gh = {}
            for r in range(H.H_FIRST, hlast + 4):
                for c in range(1, H.LASTC + 1):
                    v = ho.getCellByPosition(c - 1, r - 1)
                    if v.String and not v.getFormula().startswith("="):
                        gh[(r, c)] = v.String
            gh[(3, 1)] = ho.getCellByPosition(0, 2).String
            if gh != hcells:
                diff = [(k, gh.get(k), hcells.get(k)) for k in sorted(set(gh) | set(hcells)) if gh.get(k) != hcells.get(k)]
                errs.append(f"配布用のセルが {len(diff)} 個ちがう: {diff[:4]}")
            ghm = set()
            for r in range(H.H_FIRST, hlast + 4):
                c = 1
                while c <= H.LASTC:
                    cur = ho.createCursorByRange(ho.getCellByPosition(c - 1, r - 1))
                    cur.collapseToMergedArea()
                    a = cur.getRangeAddress()
                    if (a.EndColumn > a.StartColumn or a.EndRow > a.StartRow) and a.StartRow == r - 1 and a.StartColumn == c - 1:
                        ghm.add(f"{openpyxl.utils.get_column_letter(a.StartColumn + 1)}{a.StartRow + 1}:"
                                f"{openpyxl.utils.get_column_letter(a.EndColumn + 1)}{a.EndRow + 1}")
                    c += 1
            if ghm != hmerges:
                errs.append(f"配布用の結合ちがい: VBAのみ {sorted(ghm - hmerges)[:4]} / Pythonのみ {sorted(hmerges - ghm)[:4]}")
            # 非表示行（科目を反映）
            hidden = [r for r in range(B.FIRST, B.LAST + 1) if not plan.getRows().getByIndex(r - 1).IsVisible]
            status = "OK" if not errs else "NG"
            fails += bool(errs)
            print(f"[{status}] {case['name']}: {j['pattern']} {j['tier']}  配布用{hlast - H.H_FIRST + 1}行  セル{len(got)} 結合{len(gm)} 非表示行{len(hidden)} 教材{len(gmats)}")
            for e in errs:
                print("    ", e)
        # ダブルクリックの代わりに、☑ を付けて進捗の数式を確認
        plan.getCellRangeByName("C9").String = "☑" + plan.getCellRangeByName("C9").String
        doc.calculateAll()
        print("進捗 S8 =", repr(plan.getCellRangeByName("S8").String), " 成績表 C22 =", off.cell(doc, "成績表", "C22"))
        off.run(doc, "modReport.PrepareReport")
        print("HasChecks →", off.run(doc, "modCurriculum.HasChecks"))
        out = os.path.join(os.path.dirname(os.path.abspath(path)), "_lo_after_test.pdf")
        doc.close(True)
    return fails


if __name__ == "__main__":
    sys.exit(main(sys.argv[1]))
