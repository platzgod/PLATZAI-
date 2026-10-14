"""配布用カリキュラム（生徒に渡す A3 横の1枚）。

管理表から、チェック欄・進捗・空欄を除いて組み直す。VBA の MakeHandout（modHandout.bas）と同じ規則。
  - 科目ごとに、中身のある行だけを並べる：映像講座／受講回／塾の教材／家庭学習／共通テスト
  - 受講回の ☑ と ※ は表示しない。教材の【塾】【家】は行見出しにあるので外す
  - 共通テスト講座は、9月〜12月の位置に「講座名：回 → 回 …」の1本の帯で入れる
"""
from copy import copy

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

import data as D

SHEET = "配布用カリキュラム"
H_FIRST = 7          # データの先頭行
C0 = 3               # 1月の列（C）
NPER = D.NPER
LASTC = C0 + NPER - 1  # S
KY_START, KY_END = 11, 14  # 9月〜12月（期間の番号）
FONT = "游ゴシック"
NAVY = "17358F"
INK = "1D2433"
GRAY = "6B7488"
LINE = "C9CED8"
EMPTY = "F4F5F8"   # 中身のないマス
KIND_LABEL = {"講座": "映像講座", "回": "受講回", "塾": "塾の教材", "家": "家庭学習", "共テ": "共通テスト"}
ROW_H = {"講座": 20, "回": 14, "塾": 18, "家": 18, "共テ": 18}
CHECK = "☑"


def tint(rgb, t):
    r, g, b = (int(rgb[i:i + 2], 16) for i in (0, 2, 4))
    f = lambda x: round(x + (255 - x) * t)  # noqa: E731
    return "%02X%02X%02X" % (f(r), f(g), f(b))


def _font(sz, b=False, color=INK):
    return Font(name=FONT, sz=sz, b=b, color=color)


def _fill(rgb):
    return PatternFill("solid", fgColor=rgb)


def clean(kind, text):
    """表示用に整える。空なら ''"""
    t = (text or "").strip()
    if not t:
        return ""
    if kind == "回":
        if t.startswith(CHECK):
            t = t[1:]
        return "" if t == "※" else t
    if kind in ("塾", "家"):
        if t.startswith("（"):
            return ""
        for p in ("【塾】", "【家】"):
            if t.startswith(p):
                t = t[len(p):]
    return t


def handout_rows(ws_plan, subjects, first, block, col_ky):
    """管理表（openpyxl）→ [(科目, 種別, [(開始, 終了, 文言)])]"""
    spans_at = {}
    for m in ws_plan.merged_cells.ranges:
        spans_at[(m.min_row, m.min_col)] = m.max_col
    rows = []
    for s, subj in enumerate(D.SUBJECT_NAMES):
        ko, ky = subjects[subj]
        r0 = first + s * block
        if ko == "○":
            for off, kind in ((0, "講座"), (1, "回"), (2, "塾"), (4, "家")):
                spans = []
                c = 2
                while c <= 1 + NPER:
                    end = spans_at.get((r0 + off, c), c)
                    t = clean(kind, ws_plan.cell(r0 + off, c).value)
                    if t:
                        spans.append((c - 2, end - 2, t))
                    c = end + 1
                if spans:
                    rows.append((subj, kind, spans))
        if ky == "○":
            for k in range(block):
                name = ws_plan.cell(r0 + k, col_ky).value
                if not name:
                    continue
                seq = [clean("回", ws_plan.cell(r0 + k, col_ky + i).value) for i in range(1, 5)]
                seq = [x for x in seq if x]
                text = str(name) + ("：" + " → ".join(seq) if seq else "")
                rows.append((subj, "共テ", [(KY_START, KY_END, text)]))
    return rows


def summary_text(j, checks):
    ko = [it for it in D.ITEM_NAMES if checks[it][0]]
    ky = [it for it in D.ITEM_NAMES if checks[it][1]]
    kb = "国公立大" if j["system"].startswith("国立") else "私立大"
    return (f"{kb}（{j['bunri']}）・{j['pre']}レベル　｜　個別試験：{'・'.join(ko) or 'なし'}"
            f"　｜　共通テスト：{'・'.join(ky) or 'なし'}")


# ------------------------------------------------------------------ シートの土台（見出し）
def build_static(ws):
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = "5CAC4A"
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 8.5
    for i in range(NPER):
        ws.column_dimensions[L(C0 + i)].width = 9.8
    navy = _fill(NAVY)
    ws.merge_cells("A1:D2")
    ws["A1"] = "合格一貫カリキュラム"
    ws["A1"].font, ws["A1"].fill = _font(17, True, "FFFFFF"), navy
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    cells = [("E1:J1", '="氏名　"&IF(kName="","　　　　　　",kName)&"　さん"', 13),
             ("E2:J2", '=IF(kHighschool="","",kHighschool&"　")&kGrade', 10),
             ("K1:S1", '="第1志望　"&IF(kSchool1="","（未入力）",kSchool1)', 13),
             ("K2:S2", '=IF(kSchool2="","","第2志望　"&kSchool2&"　")&IF(kSchool3="","","第3志望　"&kSchool3)', 10)]
    for rng, f, sz in cells:
        ws.merge_cells(rng)
        c = ws[rng.split(":")[0]]
        c.value = f
        c.font = _font(sz, True, NAVY)
        c.alignment = Alignment(horizontal="left", vertical="center", indent=1, shrink_to_fit=True)
        for row in ws[rng]:
            for x in row:
                x.fill = _fill("EEF2FB")
                x.border = Border(bottom=Side(style="thin", color=LINE))
    ws.row_dimensions[1].height = 26
    ws.row_dimensions[2].height = 20
    ws.merge_cells(f"A3:{L(LASTC)}3")
    ws["A3"].font = _font(9.5, True, INK)
    ws["A3"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[3].height = 22
    # 学年の帯
    ws.row_dimensions[4].height = 13
    for rng, txt, col in ((f"{L(C0)}4:{L(C0 + 2)}4", "◀ 高2 ▶", "3462B0"), (f"{L(C0 + 3)}4:{L(LASTC)}4", "◀ 高3 ▶", NAVY)):
        ws.merge_cells(rng)
        c = ws[rng.split(":")[0]]
        c.value = txt
        c.font, c.fill = _font(8, True, "FFFFFF"), _fill(col)
        c.alignment = Alignment(horizontal="center", vertical="center")
    # 学習目的
    phases = [(0, 8, "入試基礎力を強化する", "4678C4"), (9, 10, "入試基礎力を完成させる", "3462B0"),
              (11, 14, "合格得点力を養成する（共通テスト対策スタート）", "2850A0"), (15, 16, "合格得点力を完成させる", NAVY)]
    ws.merge_cells("A5:B5")
    ws["A5"] = "学習目的"
    ws["A5"].font = _font(8.5, True, NAVY)
    ws["A5"].alignment = Alignment(horizontal="center", vertical="center")
    for a, b, txt, col in phases:
        rng = f"{L(C0 + a)}5:{L(C0 + b)}5"
        ws.merge_cells(rng)
        c = ws[rng.split(":")[0]]
        c.value = txt
        c.font, c.fill = _font(9, True, "FFFFFF"), _fill(col)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[5].height = 24
    ws.merge_cells("A6:B6")
    ws["A6"] = "科目"
    ws["A6"].font = _font(9, True, NAVY)
    ws["A6"].alignment = Alignment(horizontal="center", vertical="center")
    for i, (p, _, _) in enumerate(D.PERIODS):
        c = ws.cell(6, C0 + i, p)
        hl = p in D.HIGHLIGHT_PERIODS
        c.font = _font(9, True, INK if hl else NAVY)
        c.fill = _fill("FFD678" if hl else "EEF2FB")
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = Border(left=Side(style="thin", color=LINE), bottom=Side(style="medium", color=NAVY))
    ws.row_dimensions[6].height = 18
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = 8
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 1
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.4
    ws.print_options.horizontalCentered = True


# ------------------------------------------------------------------ 中身（VBA の MakeHandout と同じ）
def write_rows(ws, rows, summary, colors):
    for m in list(ws.merged_cells.ranges):
        if m.min_row >= H_FIRST:
            ws.unmerge_cells(str(m))
    for r in range(H_FIRST, ws.max_row + 1):
        for c in range(1, LASTC + 1):
            cell = ws.cell(r, c)
            cell.value = None
            cell.style = "Normal"
        ws.row_dimensions[r].height = None
    ws["A3"] = summary
    thin = Side(style="thin", color=LINE)
    med = Side(style="medium", color=NAVY)
    r = H_FIRST
    i = 0
    while i < len(rows):
        subj = rows[i][0]
        block = [x for x in rows if x[0] == subj]
        color = colors[subj]
        r_start = r
        for _, kind, spans in block:
            ws.row_dimensions[r].height = ROW_H[kind]
            lab = ws.cell(r, 2, KIND_LABEL[kind])
            lab.font = _font(8, True, color)
            lab.fill = _fill(tint(color, 0.88))
            lab.alignment = Alignment(horizontal="center", vertical="center")
            for c in range(2, LASTC + 1):
                ws.cell(r, c).border = Border(left=thin, right=thin, top=thin, bottom=thin)
                if c >= C0:
                    ws.cell(r, c).fill = _fill(EMPTY)
            for a, b, text in spans:
                c1, c2 = C0 + a, C0 + b
                if c2 > c1:
                    ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
                cell = ws.cell(r, c1, text)
                style_span(cell, kind, color)
                for c in range(c1, c2 + 1):
                    if kind != "回":
                        ws.cell(r, c).fill = copy(cell.fill)
            r += 1
        if r - 1 > r_start:
            ws.merge_cells(start_row=r_start, start_column=1, end_row=r - 1, end_column=1)
        a = ws.cell(r_start, 1, subj)
        a.font = _font(12 if len(subj) <= 3 else 10, True, "FFFFFF")
        a.fill = _fill(color)
        a.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for rr in range(r_start, r):
            ws.cell(rr, 1).fill = _fill(color)
            ws.cell(rr, 1).border = Border(left=med)
        for c in range(1, LASTC + 1):
            cell = ws.cell(r - 1, c)
            cell.border = Border(left=cell.border.left, right=cell.border.right, top=cell.border.top, bottom=med)
        for rr in range(r_start, r):   # 前期と後期の境（夏期Ⅱ期 | 9月）
            cell = ws.cell(rr, C0 + KY_START)
            cell.border = Border(left=med, right=cell.border.right, top=cell.border.top, bottom=cell.border.bottom)
        i += len(block)
    last = r - 1
    foot1, foot2 = last + 2, last + 3
    ws.merge_cells(start_row=foot1, start_column=1, end_row=foot1, end_column=LASTC)
    ws.cell(foot1, 1, "※ 受講回・教材の進め方は目安です。学習の進み具合に合わせて講師が調整します。"
                      "共通テスト対策は9月スタート。直前期は志望校の過去問演習が中心になります。")
    ws.cell(foot1, 1).font = _font(8.5, False, GRAY)
    ws.merge_cells(start_row=foot2, start_column=1, end_row=foot2, end_column=LASTC - 3)
    ws.cell(foot2, 1, "高校・大学受験専門 個人指導塾PLATZ ／ 〒373-0817 群馬県太田市飯塚町1417-1 ／ TEL 0276-55-4230 ／ 定休日：日・水")
    ws.cell(foot2, 1).font = _font(8.5, False, GRAY)
    ws.merge_cells(start_row=foot2, start_column=LASTC - 2, end_row=foot2, end_column=LASTC)
    ws.cell(foot2, LASTC - 2, '="作成日："&TEXT(基準日,"yyyy年m月d日")')
    ws.cell(foot2, LASTC - 2).font = _font(8.5, False, GRAY)
    ws.cell(foot2, LASTC - 2).alignment = Alignment(horizontal="right")
    ws.print_area = f"A1:{L(LASTC)}{foot2}"
    return last


def style_span(cell, kind, color):
    if kind == "講座":
        cell.font, cell.fill = _font(9, True, "FFFFFF"), _fill(color)
    elif kind == "回":
        cell.font, cell.fill = _font(8, False, INK), _fill(tint(color, 0.9))
    elif kind == "塾":
        cell.font, cell.fill = _font(8.5, True, INK), _fill(tint(color, 0.7))
    elif kind == "家":
        cell.font, cell.fill = _font(8.5, True, INK), _fill("FCEBD6")
    else:
        cell.font, cell.fill = _font(8.5, True, NAVY), _fill(tint(color, 0.8))
    cell.alignment = Alignment(horizontal="center", vertical="center", shrink_to_fit=True)
