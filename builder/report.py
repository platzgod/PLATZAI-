"""学習達成度レポート（成績表・A4縦1枚）。

  ① 志望校までの達成度 … ドーナツ4つ（総合・講座の受講・模試・家庭学習）＋大きな％
  ② 講座の受講状況     … 科目別：受講回・受講率バー・予定どおり/遅れ
  ③ 模試の偏差値の推移 … 折れ線（総合・英語・数学・国語）＋ 目標・最新・目標との差
  ④ 家庭学習推奨教材   … 周回（現在／目標）・進捗率バー

グラフ・集計の元データは印刷範囲の外（N列〜）に数式で置く。マクロなしでも更新される。
VBA の PrepareReport は ② と ④ の使っていない行を隠す（行番号は下の定数と合わせる）。
"""
from openpyxl.chart import DoughnutChart, LineChart, Reference, Series
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.chart.marker import Marker
from openpyxl.chart.series import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText
from openpyxl.drawing.text import RichTextProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
from openpyxl.drawing.text import CharacterProperties, Font as DFont, Paragraph, ParagraphProperties
from openpyxl.formatting.rule import DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

import data as D

SHEET = "成績表"
FONT = "游ゴシック"
NAVY = "17358F"
SECTION = "1F4E9E"
PALE = "EEF2FB"
INK = "1D2433"
GRAY = "6B7488"
LINE = "D5DAE3"
GRID = "E3E6EC"
TRACK = "E6E9EF"
BLUE, BLUE2, ORANGE, AQUA, YELLOW = "2A78D6", "4A90DA", "EB6834", "1BAF7A", "EDA100"
PLAN = "合格一貫カリキュラム"
MOCK = "模試・自学"

# 行の配置（VBA の PrepareReport と合わせる）
SUBJ_FIRST = 22            # ② 科目の先頭行（12科目）
MAT_FIRST_ROW = 54         # ④ 教材の先頭行（10行）
MAT_N = 10
HC = 14                    # 補助データの先頭列（N）
LASTROW = 65


def _font(sz, b=False, color=INK):
    return Font(name=FONT, sz=sz, b=b, color=color)


def _fill(c):
    return PatternFill("solid", fgColor=c)


def _txt(sz=800, color=GRAY, flat=False):
    cp = CharacterProperties(sz=sz, solidFill=color, latin=DFont(typeface=FONT), ea=DFont(typeface=FONT))
    body = RichTextProperties(rot=0, vert="horz") if flat else None
    return RichText(bodyPr=body, p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])


def _place(ws, chart, c1, r1, c2, r2):
    chart.anchor = TwoCellAnchor(_from=AnchorMarker(col=c1, row=r1), to=AnchorMarker(col=c2, row=r2))
    chart.graphical_properties = GraphicalProperties(ln=LineProperties(noFill=True))
    ws.add_chart(chart)


def _merge(ws, rng, value=None, font=None, fill=None, align=None):
    ws.merge_cells(rng)
    c = ws[rng.split(":")[0]]
    if value is not None:
        c.value = value
    if font:
        c.font = font
    if fill:
        for row in ws[rng]:
            for x in row:
                x.fill = fill
    c.alignment = align or Alignment(horizontal="center", vertical="center")
    return c


def _section(ws, r, text):
    _merge(ws, f"A{r}:L{r}", text, _font(11, True, "FFFFFF"), _fill(SECTION),
           Alignment(horizontal="left", vertical="center", indent=1))
    ws.row_dimensions[r].height = 22


def _hline(ws, r, side):
    for c in range(1, 13):
        cell = ws.cell(r, c)
        cell.border = Border(left=cell.border.left, right=cell.border.right, top=cell.border.top, bottom=side)


def build(ws):
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = "2A78D6"
    for i in range(1, 13):
        ws.column_dimensions[chr(64 + i)].width = 7.8
    ws.column_dimensions["M"].width = 3
    for i in range(HC, HC + 6):
        ws.column_dimensions[chr(64 + i)].width = 10
    thin = Side(style="thin", color=LINE)
    navy_med = Side(style="medium", color=NAVY)

    # ------------------------------------------------------------ 見出し
    _merge(ws, "A1:H1", "学習達成度レポート", _font(20, True, NAVY), align=Alignment(horizontal="left", vertical="center"))
    _merge(ws, "I1:L1", "個人指導塾 PLATZ", _font(12, True, NAVY), align=Alignment(horizontal="right", vertical="center"))
    ws.row_dimensions[1].height = 36
    _merge(ws, "I2:L2", '="作成日："&TEXT(基準日,"yyyy年m月d日")', _font(9, False, GRAY),
           align=Alignment(horizontal="right", vertical="center"))
    _hline(ws, 2, navy_med)
    ws.row_dimensions[2].height = 16
    ws.row_dimensions[3].height = 10
    info = [
        (4, [("A4:B4", "氏名"), ("C4:F4", '=IF(kName="","（未入力）",kName&"　さん")'),
             ("G4:H4", "高校・学年"), ("I4:L4", '=IF(kHighschool="","",kHighschool&"　")&kGrade')]),
        (5, [("A5:B5", "志望校"), ("C5:L5", '="第1志望　"&IF(kSchool1="","（未入力）",kSchool1)'
                                           '&IF(kSchool2="","","　／　第2志望　"&kSchool2)&IF(kSchool3="","","　／　第3志望　"&kSchool3)')]),
        (6, [("A6:B6", "コース"), ("C6:F6", '=kPatternKey'),
             ("G6:H6", "担当"), ("I6:L6", '=IF(kTantou="","",kTantou)')]),
    ]
    for r, cells in info:
        ws.row_dimensions[r].height = 22
        for i, (rng, v) in enumerate(cells):
            label = i % 2 == 0
            c = _merge(ws, rng, v, _font(9.5 if label else 11, True, NAVY if label else INK),
                       _fill(PALE) if label else None,
                       align=Alignment(horizontal="center" if label else "left", vertical="center", indent=0 if label else 1,
                                 shrink_to_fit=True))
            for row in ws[rng]:
                for x in row:
                    x.border = Border(top=thin, bottom=thin)
    ws.row_dimensions[7].height = 14

    # ------------------------------------------------------------ 補助データ（N列〜・印刷範囲外）
    ws.cell(3, HC, "集計・グラフ用（自動計算・編集不要）").font = _font(8, True, GRAY)
    # ② の科目別（N=済, O=全, P=遅れ）
    for col, t in zip(range(HC, HC + 3), ("済", "全", "遅れ")):
        ws.cell(SUBJ_FIRST - 1, col, t).font = _font(8, True, GRAY)
    end = chr(64 + 1 + D.NPER)
    for i in range(len(D.SUBJECT_NAMES)):
        r = SUBJ_FIRST + i
        k = 8 + i * 6 + 1
        rng = f"{PLAN}!$B${k}:${end}${k}"
        on = f'(設定!$B${13 + i}="○")'
        ws.cell(r, HC, f'=((COUNTIF({rng},"☑*")+COUNTIF({rng},"☑*・*")))*{on}')
        ws.cell(r, HC + 1, f'=((COUNTIF({rng},"?*")-COUNTIF({rng},"※")+COUNTIF({rng},"*・*")))*{on}')
        ws.cell(r, HC + 2, f'=(SUMPRODUCT((LEFT({rng},1)<>"☑")*({rng}<>"")*({rng}<>"※")*(期間終了日<基準日)'
                           f'*(1+ISNUMBER(FIND("・",{rng})))))*{on}')
    s_lo, s_hi = SUBJ_FIRST, SUBJ_FIRST + len(D.SUBJECT_NAMES) - 1
    N, O = chr(64 + HC), chr(64 + HC + 1)
    # ① 達成度：Q=元の値, O=グラフ用(0〜1), P=残り
    gauges = [
        ("総合達成度", f"=IFERROR(AVERAGE({chr(64 + HC + 3)}10:{chr(64 + HC + 3)}12),\"\")"),
        ("講座の受講", f'=IF(SUM({O}{s_lo}:{O}{s_hi})=0,"",SUM({N}{s_lo}:{N}{s_hi})/SUM({O}{s_lo}:{O}{s_hi}))'),
        ("模試（偏差値）", f"={MOCK}!J5"),
        ("家庭学習教材", f"={MOCK}!J7"),
    ]
    Q = chr(64 + HC + 3)
    for i, (lab, f) in enumerate(gauges):
        r = 9 + i
        ws.cell(r, HC, lab)
        ws.cell(r, HC + 3, f)
        ws.cell(r, HC + 1, f"=IF(ISNUMBER({Q}{r}),MAX(0,MIN(1,{Q}{r})),0)")
        ws.cell(r, HC + 2, f"=1-{O}{r}")
    for r in range(3, s_hi + 1):
        for c in range(HC, HC + 4):
            ws.cell(r, c).font = _font(8, False, GRAY)

    # ------------------------------------------------------------ ① 達成度
    _section(ws, 8, "① 志望校までの達成度")
    for i, (lab, _) in enumerate(gauges):
        r = 9 + i
        ch = DoughnutChart(holeSize=62)
        ch.add_data(Reference(ws, min_col=HC + 1, max_col=HC + 2, min_row=r, max_row=r), from_rows=True)
        s = ch.series[0]
        for idx, colr in ((0, BLUE if i == 0 else BLUE2), (1, TRACK)):
            pt = DataPoint(idx=idx)
            pt.graphicalProperties = GraphicalProperties(solidFill=colr, ln=LineProperties(solidFill="FFFFFF", w=19050))
            s.dPt.append(pt)
        ch.legend = None
        ch.layout = Layout(manualLayout=ManualLayout(x=0.15, y=0.06, w=0.7, h=0.9))
        _place(ws, ch, 3 * i, 8, 3 * i + 3, 15)
        c1, c3 = chr(65 + 3 * i), chr(67 + 3 * i)
        _merge(ws, f"{c1}16:{c3}16", f'=IF(ISNUMBER({Q}{r}),TEXT({Q}{r},"0%"),"—")',
               _font(20 if i == 0 else 16, True, NAVY if i == 0 else INK))
        _merge(ws, f"{c1}17:{c3}17", lab, _font(9.5, True, GRAY))
    for r in range(9, 16):
        ws.row_dimensions[r].height = 15
    ws.row_dimensions[16].height = 28
    ws.row_dimensions[17].height = 16
    _merge(ws, "A18:L18", "※ 総合達成度＝講座の受講・模試・家庭学習教材の平均（未入力の項目は除く）。"
                          "模試は偏差値40を0%、志望校の目標偏差値を100%として計算。",
           _font(7.5, False, GRAY), align=Alignment(horizontal="left", vertical="center"))
    ws.row_dimensions[19].height = 14

    # ------------------------------------------------------------ ② 講座の受講状況
    _section(ws, 20, "② 講座の受講状況（個別試験対策の映像講座）")
    for rng, t in (("A21:B21", "科目"), ("C21:D21", "受講回"), ("E21:I21", "受講率"), ("J21:L21", "予定との比較")):
        _merge(ws, rng, t, _font(9, True, NAVY), _fill(PALE))
    ws.row_dimensions[21].height = 18
    colors = dict(D.SUBJECTS)
    for i, subj in enumerate(D.SUBJECT_NAMES):
        r = SUBJ_FIRST + i
        n, o, p = (f"{chr(64 + HC + j)}{r}" for j in range(3))
        _merge(ws, f"A{r}:B{r}", subj, _font(10, True, "FFFFFF"), _fill(colors[subj]))
        _merge(ws, f"C{r}:D{r}", f'=IF({o}=0,"",{n}&" / "&{o}&" 回")', _font(10, False, INK))
        _merge(ws, f"E{r}:I{r}", f'=IF({o}=0,"",{n}/{o})', _font(9.5, True, INK),
               align=Alignment(horizontal="left", vertical="center", indent=1))
        ws[f"E{r}"].number_format = "0%"
        _merge(ws, f"J{r}:L{r}", f'=IF({o}=0,"",IF({p}=0,"予定どおり","遅れ "&{p}&" 回"))', _font(10, True, INK))
        for c in range(1, 13):
            ws.cell(r, c).border = Border(bottom=thin)
        ws.row_dimensions[r].height = 20
    ws.conditional_formatting.add(f"E{s_lo}:E{s_hi}", DataBarRule(start_type="num", start_value=0, end_type="num",
                                                                   end_value=1, color="78A0DC"))
    ws.conditional_formatting.add(f"J{s_lo}:J{s_hi}", FormulaRule(formula=[f'LEFT($J{s_lo},2)="遅れ"'],
                                                                   font=Font(color="C81E1E", b=True), stopIfTrue=True))
    ws.conditional_formatting.add(f"J{s_lo}:J{s_hi}", FormulaRule(formula=[f'$J{s_lo}="予定どおり"'],
                                                                   font=Font(color="1E8A4A", b=True)))
    ws.row_dimensions[s_hi + 1].height = 14          # 34

    # ------------------------------------------------------------ ③ 模試
    _section(ws, 35, "③ 模試の偏差値の推移")
    mock = ws.parent[MOCK]
    lc = LineChart()
    for col, colr, title in (("H", BLUE, "総合"), ("C", ORANGE, "英語"), ("D", AQUA, "数学"), ("E", YELLOW, "国語")):
        ci = ord(col) - 64
        s = Series(Reference(mock, min_col=ci, min_row=6, max_row=13), title=title)
        s.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill=colr, w=28575 if col == "H" else 19050))
        s.marker = Marker(symbol="circle", size=7)
        s.marker.graphicalProperties = GraphicalProperties(solidFill=colr, ln=LineProperties(solidFill="FFFFFF", w=9525))
        s.smooth = False
        lc.series.append(s)
    lc.set_categories(Reference(mock, min_col=1, min_row=6, max_row=13))
    lc.display_blanks = "gap"
    for ax, grid in ((lc.x_axis, False), (lc.y_axis, True)):
        ax.delete = False
        ax.txPr = _txt(750, GRAY, flat=not grid)
        ax.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill="C9CED8", w=6350))
        if grid:
            ax.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill=GRID, w=6350))
        else:
            ax.majorGridlines = None
    lc.y_axis.numFmt = "0"
    lc.legend.position = "t"
    lc.legend.txPr = _txt(850, INK)
    _place(ws, lc, 0, 35, 12, 46)
    for r in range(36, 47):
        ws.row_dimensions[r].height = 17
    # 目標・最新・差
    heads = [("A47:C47", ""), ("D47:E47", "英語"), ("F47:G47", "数学"), ("H47:I47", "国語"), ("J47", "地歴公民"),
             ("K47", "理科"), ("L47", "総合")]
    mcols = ["C", "D", "E", "F", "G", "H"]
    for rng, t in heads:
        if ":" in rng:
            _merge(ws, rng, t, _font(9, True, NAVY), _fill(PALE))
        else:
            c = ws[rng]
            c.value, c.font, c.fill = t, _font(9, True, NAVY), _fill(PALE)
            c.alignment = Alignment(horizontal="center", vertical="center")
    rows = [(48, "志望校の目標", 5, False), (49, "最新の結果", 14, False), (50, "目標との差", 15, True)]
    for r, lab, mrow, diff in rows:
        _merge(ws, f"A{r}:C{r}", lab, _font(9.5, True, NAVY))
        for (rng47, _), mc in zip(heads[1:], mcols):
            rng = rng47.replace("47", str(r))
            f = f'=IF({MOCK}!{mc}{mrow}="","—",{MOCK}!{mc}{mrow})'
            if ":" in rng:
                c = _merge(ws, rng, f, _font(10.5, True, INK))
            else:
                c = ws[rng]
                c.value, c.font = f, _font(10.5, True, INK)
                c.alignment = Alignment(horizontal="center", vertical="center")
            c.number_format = '+0.0;-0.0;0.0' if diff else "0.0"
        for c in range(1, 13):
            ws.cell(r, c).border = Border(bottom=thin)
        ws.row_dimensions[r].height = 19
    for c in range(1, 13):
        ws.cell(47, c).border = Border(bottom=thin)
    ws.row_dimensions[47].height = 18
    ws.conditional_formatting.add("D50:L50", FormulaRule(formula=['AND(ISNUMBER(D50),D50<0)'], font=Font(color="C81E1E")))
    ws.conditional_formatting.add("D50:L50", FormulaRule(formula=['AND(ISNUMBER(D50),D50>=0)'], font=Font(color="1E8A4A")))
    ws.row_dimensions[51].height = 14

    # ------------------------------------------------------------ ④ 家庭学習推奨教材
    _section(ws, 52, "④ 家庭学習推奨教材")
    for rng, t in (("A53:B53", "科目"), ("C53:H53", "教材名"), ("I53:J53", "周回（現在／目標）"), ("K53:L53", "進捗率")):
        _merge(ws, rng, t, _font(9, True, NAVY), _fill(PALE))
    ws.row_dimensions[53].height = 18
    for i in range(MAT_N):
        r, m = MAT_FIRST_ROW + i, 20 + i
        _merge(ws, f"A{r}:B{r}", f'=IF({MOCK}!A{m}="","",{MOCK}!B{m})', _font(9.5, True, GRAY))
        _merge(ws, f"C{r}:H{r}", f'=IF({MOCK}!A{m}="","",{MOCK}!A{m})', _font(10, False, INK),
               align=Alignment(horizontal="left", vertical="center", indent=1, shrink_to_fit=True))
        _merge(ws, f"I{r}:J{r}", f'=IF({MOCK}!A{m}="","",IF({MOCK}!D{m}="","-",{MOCK}!D{m})&" / "'
                                 f'&IF({MOCK}!C{m}="","-",{MOCK}!C{m})&" 周")', _font(10, False, INK))
        c = _merge(ws, f"K{r}:L{r}", f"={MOCK}!E{m}", _font(9.5, True, INK),
                   align=Alignment(horizontal="left", vertical="center", indent=1))
        c.number_format = "0%"
        for cc in range(1, 13):
            ws.cell(r, cc).border = Border(bottom=thin)
        ws.row_dimensions[r].height = 19
    ws.conditional_formatting.add(f"K{MAT_FIRST_ROW}:K{MAT_FIRST_ROW + MAT_N - 1}",
                                  DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="78A0DC"))
    ws.row_dimensions[64].height = 14
    _merge(ws, f"A{LASTROW}:L{LASTROW}",
           "高校・大学受験専門 個人指導塾PLATZ ／ 〒373-0817 群馬県太田市飯塚町1417-1 ／ TEL 0276-55-4230 ／ 定休日：日・水",
           _font(8, False, GRAY))
    ws.print_area = f"A1:L{LASTROW}"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = 9
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 1
    ws.page_margins.left = ws.page_margins.right = 0.5
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    ws.print_options.horizontalCentered = True
