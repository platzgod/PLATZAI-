"""添削結果 JSON から「<元ファイル名>_添削完了.docx」を生成する。

使い方:
    python tools/make_docx.py <添削結果.json> <出力先.docx>

仕上がり: A3 横。タイトル・生徒情報・総合評価（A〜E）の下に 3 列の表
    1. 原文（下線＝書き直す箇所）
    2. 番号ごとの直し方（〔 〕は自分の言葉で埋める）
    3. 修正後の文章（丸数字は直した箇所）＋段落の入れ替え・提出前の確認・
       アドミッションポリシー・文字起こしの確認

JSON の形式は tools/sample_result.json を参照。
- 指摘箇所は corrections[].target（原文中の文字列そのまま）で指定する。
- 丸数字は原文での出現順に自動で振るので、corrections の並び順は問わない。
- 各指摘には id を付け、他の欄では [[id]] と書くとその指摘の丸数字に置き換わる。
  revised_text の中では [[id]] が「直した箇所」の赤い丸数字になる。
"""

import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Twips

import desktop
import note

MINCHO = "游明朝"
GOTHIC = "Yu Gothic"
NAVY = RGBColor(0x1F, 0x4E, 0x79)
RED = RGBColor(0xC0, 0x00, 0x00)
BLACK = RGBColor(0x00, 0x00, 0x00)
GRAY = RGBColor(0x59, 0x59, 0x59)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
COL_WIDTHS = (6700, 7700, 7720)  # twips。A3 横の本文幅いっぱい
MARKER = re.compile(r"\[\[([^\[\]]+)\]\]")
RANKS = {
    "A": "合格水準・完成度が高い",
    "B": "合格圏・細部の改善で仕上がる",
    "C": "要改善・構成か論証に課題",
    "D": "大幅な書き直しが必要",
    "E": "設問に答えていない・未完成",
}


def circled(n):
    """1〜20 は ①〜⑳、21 以降は (21)(22)…"""
    return chr(0x2460 + n - 1) if 1 <= n <= 20 else f"({n})"


# ---------- 書式ヘルパー ----------

def run(p, text, font=MINCHO, size=10, bold=False, color=BLACK, underline=False):
    r = p.add_run(text)
    r.font.name = font
    r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), font)
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.underline = underline
    return r


def para(container, line=None, before=0, after=0, left=0, hanging=0, first=0, align=None):
    p = container.add_paragraph()
    f = p.paragraph_format
    if line:
        f.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        f.line_spacing = Twips(line)
    f.space_before = Twips(before)
    f.space_after = Twips(after)
    if left:
        f.left_indent = Twips(left)
    if hanging:
        f.first_line_indent = Twips(-hanging)
    elif first:
        f.first_line_indent = Twips(first)
    if align is not None:
        p.alignment = align
    return p


def set_cell(cell, width, fill=None):
    cell.width = Twips(width)
    if fill:
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill)
        cell._tc.get_or_add_tcPr().append(shd)


def style_table(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{edge}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "4")
        b.set(qn("w:space"), "0")
        b.set(qn("w:color"), "1F4E79")
        borders.append(b)
    tblPr.append(borders)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    mar = OxmlElement("w:tblCellMar")
    for edge, w in (("top", 120), ("left", 180), ("bottom", 120), ("right", 180)):
        m = OxmlElement(f"w:{edge}")
        m.set(qn("w:w"), str(w))
        m.set(qn("w:type"), "dxa")
        mar.append(m)
    tblPr.append(mar)


def clear_cell(cell):
    """セルに最初からある空段落を消す（中身を書き終えてから呼ぶ）。"""
    first = cell.paragraphs[0]
    if not first.text and len(cell.paragraphs) > 1:
        first._element.getparent().remove(first._element)


# ---------- 内容の組み立て ----------

def locate(original, corrections):
    """各指摘の (開始, 終了, 指摘) を出現順に返す。見つからない・重なる場合はエラー。"""
    spans = []
    for c in corrections:
        target, occurrence = c["target"], c.get("occurrence", 1)
        start = -1
        for _ in range(occurrence):
            start = original.find(target, start + 1)
            if start < 0:
                sys.exit(f"エラー: 原文に「{target}」が{occurrence}回目まで見つかりません。")
        spans.append((start, start + len(target), c))
    spans.sort(key=lambda s: s[0])
    for (_, e1, c1), (s2, _, c2) in zip(spans, spans[1:]):
        if s2 < e1:
            sys.exit(f"エラー: 指摘箇所が重なっています:「{c1['target']}」と「{c2['target']}」")
    return spans


def numbering(spans):
    ids = {}
    for n, (_, _, c) in enumerate(spans, 1):
        cid = c.get("id", str(n))
        if cid in ids:
            sys.exit(f"エラー: id「{cid}」が重複しています。")
        ids[cid] = n
    return ids


def resolve(text, ids):
    """[[id]] を丸数字の文字に置き換える。"""
    def sub(m):
        if m.group(1) not in ids:
            sys.exit(f"エラー: [[{m.group(1)}]] に対応する指摘がありません。")
        return circled(ids[m.group(1)])
    return MARKER.sub(sub, text or "")


def body_para(cell, first_indent):
    return para(cell, line=320, after=80, first=first_indent, align=WD_ALIGN_PARAGRAPH.JUSTIFY)


def fill_original(cell, original, spans):
    """原文。書き直す箇所は「丸数字＋本文」に下線、丸数字は赤の太字。"""
    pieces = []  # (文字列, 番号 or None)
    pos = 0
    for n, (start, end, _) in enumerate(spans, 1):
        if pos < start:
            pieces.append((original[pos:start], None))
        pieces.append((original[start:end], n))
        pos = end
    if pos < len(original):
        pieces.append((original[pos:], None))

    p = body_para(cell, 200)
    for text, number in pieces:
        if number is not None:
            run(p, circled(number), bold=True, color=RED, underline=True)
        for i, part in enumerate(text.split("\n")):
            if i > 0:
                p = body_para(cell, 200)
            if part:
                run(p, part, underline=number is not None)


def fill_advice(cell, spans, ids):
    """番号ごとの直し方: 太字の見出し行、理由（灰色）、型（灰色）。"""
    for n, (_, _, c) in enumerate(spans, 1):
        p = para(cell, line=300, before=0 if n == 1 else 40, after=40, left=615, hanging=300)
        run(p, circled(n), font=GOTHIC, bold=True, color=RED)
        run(p, "　" + resolve(c["fix"], ids), font=GOTHIC, bold=True)

        notes = []
        if c.get("reason") or c.get("criteria"):
            criteria = "・".join(c.get("criteria", []))
            notes.append(("理由：", (f"［{criteria}］" if criteria else "") + resolve(c.get("reason"), ids)))
        if c.get("template"):
            notes.append(("型：", resolve(c["template"], ids)))
        for i, (label, text) in enumerate(notes):
            q = para(cell, line=260, after=70 if i == len(notes) - 1 else 20, left=1095, hanging=360)
            run(q, label, size=9, bold=True, color=GRAY)
            run(q, text, size=9, color=GRAY)


def fill_revised(cell, revised, ids):
    """修正後の文章。[[id]] の位置に赤い丸数字を入れる。"""
    for line in revised.split("\n"):
        if not line.strip():
            continue
        p = body_para(cell, 180)
        pos = 0
        for m in MARKER.finditer(line):
            if pos < m.start():
                run(p, line[pos:m.start()])
            run(p, resolve(m.group(0), ids), size=9, bold=True, color=RED)
            pos = m.end()
        if pos < len(line):
            run(p, line[pos:])


def section(cell, title, lines):
    if not lines:
        return
    p = para(cell, line=300, before=80, after=40)
    run(p, title, font=GOTHIC, size=9.5, bold=True, color=NAVY)
    for line in lines:
        q = para(cell, line=270, after=40, left=390, hanging=180)
        run(q, "・" + line, size=9)


def char_count(text):
    return len(re.sub(r"\s", "", MARKER.sub("", text)))


def build(data, out_path):
    original = data["original_text"].replace("\r\n", "\n")
    revised = data["revised_text"].replace("\r\n", "\n")
    spans = locate(original, data.get("corrections", []))
    ids = numbering(spans)
    resolve(revised, ids)  # 未定義の [[id]] を先に検出する

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Twips(23814), Twips(16839)  # A3 横
    sec.top_margin, sec.bottom_margin = Twips(800), Twips(720)
    sec.left_margin = sec.right_margin = Twips(840)
    normal = doc.styles["Normal"]
    normal.font.name = MINCHO
    normal.font.size = Pt(10)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), MINCHO)
    normal.paragraph_format.space_after = Pt(0)

    # タイトルと生徒情報
    doc_type = data.get("doc_type", "小論文")
    p = para(doc, after=40, align=WD_ALIGN_PARAGRAPH.CENTER)
    p.paragraph_format.space_after = Twips(40)
    run(p, f"{doc_type}　書き直し箇所と直し方・修正後の文章", font=GOTHIC, size=18, bold=True, color=NAVY)
    school = "　".join(x for x in (data.get("university"), data.get("faculty")) if x) or "志望校不明"
    info = f"{data.get('student') or '生徒名不明'}　｜　{school}"
    if data.get("prompt"):
        info += f"　｜　設問：{data['prompt']}"
    p = para(doc, after=80, align=WD_ALIGN_PARAGRAPH.CENTER)
    run(p, info, font=GOTHIC, size=10, color=GRAY)

    # 総合評価
    overall = data.get("overall")
    if overall:
        rank = overall["rank"]
        if rank not in RANKS:
            sys.exit(f"エラー: 総合評価「{rank}」は {'・'.join(RANKS)} のいずれかにしてください。")
        p = para(doc, after=120)
        run(p, "総合評価　", font=GOTHIC, size=12, bold=True, color=NAVY)
        run(p, rank, font=GOTHIC, size=16, bold=True, color=RED)
        run(p, f"（{RANKS[rank]}）　", font=GOTHIC, size=9, color=GRAY)
        run(p, resolve(overall.get("comment"), ids), size=10.5)

    # 3 列の表
    table = doc.add_table(rows=2, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    style_table(table)
    last = circled(len(spans)) if spans else ""
    headers = (
        "1. 原文（下線＝書き直す箇所）",
        "2. 番号ごとの直し方（〔 〕は自分の言葉で埋める）",
        f"3. 修正後の文章（①〜{last}は直した箇所）　※参考例" if len(spans) > 1
        else "3. 修正後の文章（丸数字は直した箇所）　※参考例",
    )
    for i, (title, width) in enumerate(zip(headers, COL_WIDTHS)):
        head, body = table.cell(0, i), table.cell(1, i)
        set_cell(head, width, fill="1F4E79")
        set_cell(body, width)
        hp = head.paragraphs[0]
        hp.paragraph_format.space_after = Twips(120)
        run(hp, title, font=GOTHIC, size=11, bold=True, color=WHITE)

    left, middle, right = (table.cell(1, i) for i in range(3))
    fill_original(left, original, spans)
    if spans:
        fill_advice(middle, spans, ids)
    else:
        run(para(middle), "書き直しが必要な箇所はありません。")
    fill_revised(right, revised, ids)

    section(right, "【段落の入れ替え】", [resolve(x, ids) for x in data.get("reorder", [])])
    checks = [resolve(x, ids) for x in data.get("checklist", [])]
    before, after = char_count(original), char_count(revised)
    limit = data.get("char_limit")
    if limit:
        checks.append(
            f"字数：制限 {limit}字 → 原文 約{before}字（{before * 100 // limit}%）→ 修正後 約{after}字"
            f"（{after * 100 // limit}%）。番号記号と空白を除き、〔 〕を含む。"
        )
    else:
        checks.append(
            f"文字数の目安：原文 約{before}字 → 修正後 約{after}字"
            "（番号記号と空白を除く、〔 〕を含む）。字数制限に合わせて調整する。"
        )
    section(right, "【提出前の確認】", checks)

    policy = data.get("admission_policy") or {}
    ap_lines = list(policy.get("summary", []))
    ap_lines += [f"出典：{url}" for url in policy.get("sources", [])]
    if policy.get("note"):
        ap_lines.append(policy["note"])
    if not policy:
        ap_lines = ["志望校が特定できなかったため、固定基準のみで添削した。"]
    section(right, "【参照したアドミッションポリシー】", ap_lines)
    section(right, "【文字起こしの確認】", [resolve(x, ids) for x in data.get("transcription_notes", [])])

    for cell in (left, middle, right):
        clear_cell(cell)

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    unused = [cid for cid in ids if f"[[{cid}]]" not in revised]
    if unused:
        print(f"注意: 修正後の文章に番号が入っていない指摘: {', '.join(circled(ids[c]) for c in unused)}")
    flags = []
    if not data.get("prompt"):
        flags.append("設問なし")
    if not data.get("admission_policy"):
        flags.append("志望校不明")
    if data.get("transcription_notes"):
        flags.append("文字起こし要確認")
    if "〔" in revised:
        flags.append("〔 〕あり")
    note.append(Path(out_path).parent, Path(out_path).stem.removesuffix("_添削完了"), "納品",
                rank=(overall or {}).get("rank", ""), count=f"{len(spans)}件", notes="・".join(flags))
    print(f"保存しました: {out_path}（指摘 {len(spans)} 件）")
    for copied in desktop.save(out_path):
        print(f"デスクトップにも保存しました: {copied}")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    with open(sys.argv[1], encoding="utf-8") as f:
        build(json.load(f), sys.argv[2])
