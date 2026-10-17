"""小論文の添削シート（.docx）を作る道具。

    python3 tools/tensaku.py count 提出/答案.txt        # 字数を数える（空白・改行を除く）
    python3 tools/tensaku.py render 添削.json 添削結果/xxx.docx   # 添削シートを作る

添削の中身（点数・コメント・書き直し例など）は Claude が JSON で書き、
このスクリプトが字数を数え直してから A4 の Word 文書に組む。
"""
import json
import sys
from pathlib import Path

import jsonschema
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

# ワークシートと同じ配色
GREEN, ORANGE, INK, TEXT, MUTED = "2C5F2D", "D9782D", "17321A", "23301F", "5B6B54"
CARD, CARD_LINE, PEACH = "F6F8F3", "DDE8CF", "FCEFE3"
GOTHIC, MINCHO = "游ゴシック", "游明朝"

CRITERIA = ["設問への応答", "構成・論理", "根拠・具体性", "学部・アドミッションポリシーとの適合", "表現・表記"]

S = {"type": "string"}
SA = {"type": "array", "items": S}


def obj(**props):
    return {"type": "object", "properties": props, "required": list(props), "additionalProperties": False}


LABELED = obj(label=S, text=S)

TENSAKU_SCHEMA = obj(
    student=S,                    # 生徒名（空文字なら「無記名」）
    date=S,                       # 添削日 YYYY-MM-DD
    target=obj(
        university=S, faculty=S, department=S, admission=S,
        chars_min={"type": ["integer", "null"]},   # 「◯字以内」なら null
        chars_max={"type": "integer"},
        minutes={"type": ["integer", "null"]},
        topic=S,                  # 題材・テーマ
        question=S,               # 設問文（分かる範囲で）
    ),
    essay=SA,                     # 答案の原文（段落ごと。手を加えない）
    scores={"type": "array", "items": obj(name=S, score={"type": "integer"}, max={"type": "integer"}, comment=S)},
    summary=S,                    # 総評（3〜5文）
    strengths=SA,                 # 良い点（2〜4個）
    improvements={"type": "array", "items": obj(point=S, quote=S, advice=S)},   # 改善点（3〜5個）
    corrections={"type": "array", "items": obj(before=S, after=S, reason=S)},   # 表記・表現の直し
    structure=obj(
        student={"type": "array", "items": LABELED},   # 答案の段落ごとの中身（要約）
        advice={"type": "array", "items": LABELED},    # おすすめの構成メモ
    ),
    rewrite={"type": "array", "items": LABELED},       # 書き直し例（段落ごと）
    ap=obj(points=SA, sources=SA, link=S),             # AP の要点・出典 URL・答案とのつながり
    next_steps=SA,                                      # 次回の課題（3個）
)


class TensakuError(ValueError):
    pass


def count_chars(text):
    """原稿用紙の字数（空白・改行を除いた文字数）。"""
    return sum(1 for ch in text if not ch.isspace())


def chars_label(t):
    return f"{t['chars_min']}〜{t['chars_max']}字" if t.get("chars_min") else f"{t['chars_max']}字以内"


def char_verdict(total, t):
    lo = t["chars_min"] or int(t["chars_max"] * 0.9)
    if total > t["chars_max"]:
        return f"字数オーバー（{total - t['chars_max']}字超過）。本番では減点・採点対象外になりうる"
    if total < int(t["chars_max"] * 0.9):
        return f"9割（{int(t['chars_max'] * 0.9)}字）に届いていない" + ("・下限未満" if total < lo else "")
    return "規定の範囲内"


def validate(d):
    try:
        jsonschema.validate(d, TENSAKU_SCHEMA)
    except jsonschema.ValidationError as e:
        raise TensakuError("/".join(map(str, e.absolute_path)) + ": " + e.message) from None
    errs = []
    if not d["essay"]:
        errs.append("essay（答案の原文）が空です")
    for s in d["scores"]:
        if not 0 <= s["score"] <= s["max"]:
            errs.append(f"{s['name']}: 得点 {s['score']} が満点 {s['max']} を超えています")
    if errs:
        raise TensakuError("\n".join(errs))


# ---------------------------------------------------------------- docx helpers
TBLPR_ORDER = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize",
               "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook",
               "tblCaption", "tblDescription"]
TCPR_ORDER = ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar",
              "textDirection", "tcFitText", "vAlign", "hideMark"]


def put(parent, el, order):
    """OOXML の決まった順番を守って子要素を入れる（同名の要素は置き換える）。"""
    name = el.tag.split("}")[1]
    old = parent.find(qn("w:" + name))
    if old is not None:
        parent.remove(old)
    later = order[order.index(name) + 1:]
    for child in parent:
        if child.tag.split("}")[1] in later:
            child.addprevious(el)
            return
    parent.append(el)


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    put(tcPr, shd, TCPR_ORDER)


def borders(table, color=CARD_LINE, size=4):
    tblPr = table._tbl.tblPr
    el = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), str(size))
        b.set(qn("w:color"), color)
        el.append(b)
    put(tblPr, el, TBLPR_ORDER)
    mar = OxmlElement("w:tblCellMar")  # セル内の余白
    for side, w in (("top", 40), ("left", 100), ("bottom", 40), ("right", 100)):
        m = OxmlElement(f"w:{side}")
        m.set(qn("w:w"), str(w))
        m.set(qn("w:type"), "dxa")
        mar.append(m)
    put(tblPr, mar, TBLPR_ORDER)


def widths(table, mm):
    """列幅を固定する（表の幅・列グリッド・各セルの3か所に書かないと Word/LibreOffice で効かない）。"""
    table.autofit = False
    tblPr = table._tbl.tblPr
    tw = OxmlElement("w:tblW")
    tw.set(qn("w:w"), str(int(sum(mm) * 56.7)))
    tw.set(qn("w:type"), "dxa")
    put(tblPr, tw, TBLPR_ORDER)
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    put(tblPr, lay, TBLPR_ORDER)
    for gc, w in zip(table._tbl.tblGrid.findall(qn("w:gridCol")), mm):
        gc.set(qn("w:w"), str(int(w * 56.7)))
    for row in table.rows:
        for cell, w in zip(row.cells, mm):
            cell.width = Mm(w)


def run(p, text, size=10.5, bold=False, color=TEXT, font=GOTHIC):
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = RGBColor.from_string(color)
    r.font.name = font
    rPr = r._r.get_or_add_rPr()
    fonts = rPr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rPr.insert(0, fonts)
    fonts.set(qn("w:eastAsia"), font)
    return r


def para(container, text="", size=10.5, bold=False, color=TEXT, font=GOTHIC, align=None, after=4, before=0):
    p = container.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(after)
    pf.space_before = Pt(before)
    pf.line_spacing = 1.25
    if align:
        p.alignment = align
    if text:
        run(p, text, size, bold, color, font)
    return p


def cell_text(cell, text, size=10, bold=False, color=TEXT, font=GOTHIC, align=None):
    cell.paragraphs[0].text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.2
    if align:
        p.alignment = align
    lines = text.split("\n") if text else [""]
    run(p, lines[0], size, bold, color, font)
    for line in lines[1:]:
        q = cell.add_paragraph()
        q.paragraph_format.space_after = Pt(0)
        q.paragraph_format.line_spacing = 1.2
        run(q, line, size, bold, color, font)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def heading(doc, text, mark="■"):
    p = para(doc, before=10, after=4)
    run(p, mark + " ", 13, True, ORANGE)
    run(p, text, 13, True, GREEN)
    return p


def bullets(doc, items, mark="・", color=TEXT):
    for t in items:
        p = para(doc, after=2)
        p.paragraph_format.left_indent = Mm(4)
        p.paragraph_format.first_line_indent = Mm(-4)
        run(p, mark, 10.5, True, color)
        run(p, t, 10.5)


def labeled_table(doc, rows, total_w=170, label_w=24, font=GOTHIC, size=10.5, counts=True, fill=CARD):
    t = doc.add_table(rows=0, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(t)
    for label, text in rows:
        c = t.add_row().cells
        lab = label + (f"\n{count_chars(text)}字" if counts else "")
        cell_text(c[0], lab, 9.5, True, GREEN, align=WD_ALIGN_PARAGRAPH.CENTER)
        shade(c[0], fill)
        cell_text(c[1], text, size, font=font)
    widths(t, [label_w, total_w - label_w])
    return t


# ---------------------------------------------------------------- build
def build(d):
    t = d["target"]
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(20)
    sec.top_margin = sec.bottom_margin = Mm(18)
    zoom = doc.settings.element.find(qn("w:zoom"))
    if zoom is not None and zoom.get(qn("w:percent")) is None:
        zoom.set(qn("w:percent"), "100")
    st = doc.styles["Normal"]
    st.font.name = GOTHIC
    st.element.rPr.rFonts.set(qn("w:eastAsia"), GOTHIC)
    st.font.size = Pt(10.5)

    # フッター
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(fp, "小論文添削（ショー）", 8, color=MUTED)

    # タイトル帯
    band = doc.add_table(rows=1, cols=2)
    band.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = band.rows[0].cells
    cell_text(c[0], "小論文 添削シート", 20, True, "FFFFFF")
    cell_text(c[1], f"{d['date']}\n{d['student'] or '無記名'}", 10.5, True, "FFFFFF", align=WD_ALIGN_PARAGRAPH.RIGHT)
    for x in c:
        shade(x, GREEN)
    widths(band, [110, 60])

    # 基本情報
    essay_total = count_chars("".join(d["essay"]))
    info = [("志望", " ".join(x for x in (t["university"], t["faculty"], t["department"], t["admission"]) if x)),
            ("題材", t["topic"]),
            ("設問", t["question"] or "（記載なし）"),
            ("字数", f"制限 {chars_label(t)}" + (f"／試験時間 {t['minutes']}分" if t.get("minutes") else "") +
             f"　→　答案 {essay_total}字（{char_verdict(essay_total, t)}）")]
    para(doc, after=2)
    tb = doc.add_table(rows=0, cols=2)
    borders(tb)
    for k, v in info:
        r = tb.add_row().cells
        cell_text(r[0], k, 9.5, True, GREEN, align=WD_ALIGN_PARAGRAPH.CENTER)
        shade(r[0], CARD)
        cell_text(r[1], v, 10)
    widths(tb, [20, 150])

    # 総合点と観点別
    total = sum(s["score"] for s in d["scores"])
    full = sum(s["max"] for s in d["scores"])
    heading(doc, "評価")
    p = para(doc, after=4)
    run(p, "総合 ", 12, True, INK)
    run(p, str(total), 28, True, ORANGE)
    run(p, f" / {full}点", 12, True, INK)
    sc = doc.add_table(rows=1, cols=3)
    borders(sc)
    for cell, h in zip(sc.rows[0].cells, ("観点", "得点", "コメント")):
        cell_text(cell, h, 9.5, True, "FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER)
        shade(cell, GREEN)
    for s in d["scores"]:
        r = sc.add_row().cells
        cell_text(r[0], s["name"], 9.5, True, INK)
        cell_text(r[1], f"{s['score']} / {s['max']}", 11, True, ORANGE, align=WD_ALIGN_PARAGRAPH.CENTER)
        cell_text(r[2], s["comment"], 9.5)
    widths(sc, [38, 20, 112])

    heading(doc, "総評")
    para(doc, d["summary"], after=4)

    heading(doc, "良かったところ", "◎")
    bullets(doc, d["strengths"], "✓ ", GREEN)

    heading(doc, "直すと伸びるところ", "▲")
    for i, im in enumerate(d["improvements"], 1):
        p = para(doc, after=1, before=3)
        run(p, f"{i}. {im['point']}", 10.5, True, INK)
        if im["quote"]:
            q = para(doc, after=1)
            q.paragraph_format.left_indent = Mm(5)
            run(q, "答案：", 9.5, True, MUTED)
            run(q, f"「{im['quote']}」", 9.5, color=MUTED, font=MINCHO)
        a = para(doc, after=2)
        a.paragraph_format.left_indent = Mm(5)
        run(a, "→ ", 10.5, True, ORANGE)
        run(a, im["advice"], 10.5)

    # 原文
    heading(doc, f"あなたの答案（{essay_total}字）")
    labels = [s["label"] for s in d["structure"]["student"]]
    rows = [((labels[i] if i < len(labels) else f"第{i + 1}段落"), text) for i, text in enumerate(d["essay"])]
    labeled_table(doc, rows, font=MINCHO)

    if d["corrections"]:
        heading(doc, "表記・表現の直し")
        ct = doc.add_table(rows=1, cols=3)
        borders(ct)
        for cell, h in zip(ct.rows[0].cells, ("答案", "直し", "理由")):
            cell_text(cell, h, 9.5, True, "FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER)
            shade(cell, GREEN)
        for co in d["corrections"]:
            r = ct.add_row().cells
            cell_text(r[0], co["before"], 9.5, font=MINCHO)
            cell_text(r[1], co["after"], 9.5, True, GREEN, font=MINCHO)
            cell_text(r[2], co["reason"], 9)
        widths(ct, [60, 60, 50])

    # 構成（ワークシートの Q3 と同じ考え方）
    heading(doc, "構成の見直し（Q3 段落メモ）")
    para(doc, "左があなたの答案の骨組み、右がおすすめの骨組み。書く前にこのメモを作れば、構成は安定する。", 9.5, color=MUTED)
    stt = doc.add_table(rows=1, cols=2)
    borders(stt)
    for cell, h in zip(stt.rows[0].cells, ("あなたの答案", "おすすめの構成メモ")):
        cell_text(cell, h, 9.5, True, "FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER)
        shade(cell, GREEN if h == "あなたの答案" else ORANGE)
    sa, sb = d["structure"]["student"], d["structure"]["advice"]
    for i in range(max(len(sa), len(sb))):
        r = stt.add_row().cells
        if i < len(sa):
            cell_text(r[0], f"【{sa[i]['label']}】{sa[i]['text']}", 9.5)
        if i < len(sb):
            cell_text(r[1], f"【{sb[i]['label']}】{sb[i]['text']}", 9.5)
            shade(r[1], PEACH)
    widths(stt, [85, 85])

    # 書き直し例
    rw_total = count_chars("".join(p["text"] for p in d["rewrite"]))
    heading(doc, f"書き直し例（{rw_total}字）")
    para(doc, "あなたの考えを生かしたまま直した一例。丸写しせず、どこが変わったかを比べて読むこと。", 9.5, color=MUTED)
    labeled_table(doc, [(p["label"], p["text"]) for p in d["rewrite"]], font=MINCHO, fill=PEACH)

    # AP
    heading(doc, "アドミッションポリシーとのつながり")
    bullets(doc, d["ap"]["points"])
    if d["ap"]["link"]:
        para(doc, d["ap"]["link"], after=2, before=2)
    if d["ap"]["sources"]:
        para(doc, "出典：" + "　".join(d["ap"]["sources"]), 8, color=MUTED)

    heading(doc, "次回の課題", "★")
    for i, s in enumerate(d["next_steps"], 1):
        p = para(doc, after=2)
        run(p, f"{i}．", 10.5, True, ORANGE)
        run(p, s, 10.5)
    return doc, essay_total, rw_total


def render(d, out):
    validate(d)
    doc, essay_total, rw_total = build(d)
    Path(out).parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    t = d["target"]
    lo = t["chars_min"] or int(t["chars_max"] * 0.9)
    rw_ok = max(lo, int(t["chars_max"] * 0.9)) <= rw_total <= t["chars_max"]
    return {"essay_chars": essay_total, "essay_verdict": char_verdict(essay_total, t),
            "rewrite_chars": rw_total, "rewrite_ok": rw_ok,
            "score": sum(s["score"] for s in d["scores"])}


def main(argv):
    if len(argv) >= 2 and argv[0] == "count":
        text = sys.stdin.read() if argv[1] == "-" else Path(argv[1]).read_text(encoding="utf-8")
        paras = [p for p in text.splitlines() if p.strip()]
        print(json.dumps({"total": count_chars(text), "paragraphs": [count_chars(p) for p in paras]},
                         ensure_ascii=False))
    elif len(argv) == 3 and argv[0] == "render":
        d = json.loads(Path(argv[1]).read_text(encoding="utf-8"))
        try:
            print(json.dumps(render(d, argv[2]), ensure_ascii=False))
        except TensakuError as e:
            sys.exit(f"添削 JSON に問題があります:\n{e}")
    else:
        sys.exit(__doc__)


if __name__ == "__main__":
    main(sys.argv[1:])
