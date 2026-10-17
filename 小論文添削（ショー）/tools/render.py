"""小論文ワークシート（.pptx）を、内容 JSON から組み立てるレンダラー。

東洋大学 食環境科学部版のデザイン（16:9・緑×オレンジ・游ゴシック）を再現する。
内容は JSON（schema.py の WORKSHEET_SCHEMA）で受け取り、レイアウトは文字量から自動で決める。

    python3 tools/render.py content.json out.pptx
"""
import copy
import json
import math
import sys
import unicodedata

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

# ---------------------------------------------------------------- design tokens
GREEN = "2C5F2D"
ORANGE = "D9782D"
INK = "17321A"        # 見出し
TEXT = "23301F"       # 本文
MUTED = "5B6B54"      # 補足・フッター
CARD = "F6F8F3"       # カード地
CARD_LINE = "DDE8CF"  # カード枠
PEACH = "FCEFE3"      # 強調カード地
WHITE = "FFFFFF"
FONT = "游ゴシック"

SLIDE_W, SLIDE_H = 13.333, 7.5
LEFT, CONTENT_W = 0.60, 12.13
FULL_RIGHT = LEFT + CONTENT_W

CIRCLED = "①②③④⑤⑥⑦⑧⑨⑩"


def rgb(hex_):
    return RGBColor.from_string(hex_)


# ---------------------------------------------------------------- text metrics
def text_units(s):
    """全角=1、半角=0.55 として文字幅を数える（行数の見積もり用）。"""
    n = 0.0
    for ch in s:
        n += 1.0 if unicodedata.east_asian_width(ch) in ("F", "W", "A") else 0.55
    return n


def est_lines(text, width_in, pt):
    # 和文は字間が入るため、1em 幅より少し少なめに見積もる
    per_line = max(1.0, (width_in - 0.1) * 72 / pt * 0.93)
    return sum(max(1, math.ceil(text_units(line) / per_line)) for line in text.split("\n"))


def est_height(paragraphs, width_in, pt, line=1.15, space_after_pt=0):
    """段落リストの描画高さ（インチ）を見積もる。"""
    lines = sum(est_lines(p, width_in, pt) for p in paragraphs)
    return lines * pt * line * 1.3 / 72 + len(paragraphs) * space_after_pt / 72 + 0.06


def fit_pt(paragraphs, width_in, height_in, max_pt, min_pt, space_after_pt=0):
    """枠に収まる最大のフォントサイズ（0.5pt 刻み）を返す。"""
    pt = max_pt
    while pt > min_pt and est_height(paragraphs, width_in, pt, space_after_pt=space_after_pt) > height_in:
        pt -= 0.5
    return pt


def count_chars(text):
    """原稿用紙の字数（空白・改行を除いた文字数）。"""
    return sum(1 for ch in text if not ch.isspace())


# ---------------------------------------------------------------- primitives
class Deck:
    def __init__(self, data):
        self.d = data
        self.prs = Presentation()
        self.prs.slide_width = Inches(SLIDE_W)
        self.prs.slide_height = Inches(SLIDE_H)
        self.blank = self.prs.slide_layouts[6]
        meta = data["meta"]
        self.footer = meta.get("footer") or f"{meta['org']}　{meta['university']} {meta['faculty']} 小論文ワークシート"

    # -- shapes
    def box(self, slide, x, y, w, h, fill=None, line=None, line_w=1.0, dash=False, radius=6000):
        shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        shp.adjustments[0] = radius / 100000
        if fill:
            shp.fill.solid()
            shp.fill.fore_color.rgb = rgb(fill)
        else:
            shp.fill.background()
        if line:
            shp.line.color.rgb = rgb(line)
            shp.line.width = Pt(line_w)
            if dash:
                shp.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        else:
            shp.line.fill.background()
        shp.shadow.inherit = False
        tf = shp.text_frame
        tf.margin_left = tf.margin_right = Inches(0.1)
        tf.margin_top = tf.margin_bottom = Inches(0.05)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        return shp

    def pill(self, slide, x, y, w, h, label, fill=GREEN, pt=14, color=WHITE, line=None, radius=30000):
        shp = self.box(slide, x, y, w, h, fill=fill, line=line, line_w=1.25, radius=radius)
        tf = shp.text_frame
        tf.margin_left = tf.margin_right = Inches(0.04)
        tf.margin_top = tf.margin_bottom = 0
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        self._run(p, label, pt, True, color)
        return shp

    def text(self, slide, x, y, w, h, paras, pt=16, color=TEXT, bold=False, align="l",
             anchor="t", space_after=0, autofit=True):
        """paras: 文字列、または段落のリスト。段落は文字列か [(text, {pt,bold,color}), ...]。"""
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.05)
        tf.margin_top = tf.margin_bottom = Inches(0.03)
        tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE}[anchor]
        if autofit:
            bp = tf._txBody.find(qn("a:bodyPr"))
            etree.SubElement(bp, qn("a:spAutoFit"))
        if isinstance(paras, str):
            paras = [paras]
        for i, para in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
            p.line_spacing = 1.15
            p.space_after = Pt(space_after if i < len(paras) - 1 else 0)
            runs = [(para, {})] if isinstance(para, str) else para
            for t, o in runs:
                self._run(p, t, o.get("pt", pt), o.get("bold", bold), o.get("color", color))
        return tb

    @staticmethod
    def _run(p, text, pt, bold, color):
        r = p.add_run()
        r.text = text
        f = r.font
        f.size = Pt(pt)
        f.bold = bold
        f.color.rgb = rgb(color)
        f.name = FONT
        rPr = r._r.get_or_add_rPr()
        ea = rPr.find(qn("a:ea"))
        if ea is None:
            ea = etree.SubElement(rPr, qn("a:ea"))
        ea.set("typeface", FONT)

    def group(self, slide, shapes):
        """作成済みの図形をグループにまとめる（クリック表示アニメーションの単位）。"""
        grp = slide.shapes.add_group_shape()
        for s in shapes:
            grp._element.append(s._element)  # spTree から group へ移動
        grp._element.recalculate_extents()
        return grp

    # -- slide furniture
    def new_slide(self, kicker, title, sub=None, badge=None):
        s = self.prs.slides.add_slide(self.blank)
        self.text(s, 0.60, 0.38, 9.50, 0.40, kicker, pt=15, color=ORANGE, bold=True)
        self.text(s, 0.60, 0.72, 10.0 if badge else 12.1, 0.80, title,
                  pt=fit_pt([title], 10.0 if badge else 12.1, 0.80, 32, 24), color=INK, bold=True)
        if sub:
            self.text(s, 0.60, 1.48, 12.00, 0.45, sub, pt=16, color=MUTED)
        if badge:
            self.pill(s, 10.75, 0.45, 2.00, 0.50, badge, fill=ORANGE, pt=16, radius=30000)
        n = len(self.prs.slides)
        self.text(s, 0.50, 7.08, 9.00, 0.30, self.footer, pt=10, color=MUTED)
        self.text(s, 12.20, 7.08, 0.65, 0.30, str(n), pt=10, color=MUTED, align="r")
        return s

    def bar(self, slide, label, body, y=6.25, h=0.62, pt=18):
        self.box(slide, LEFT, y, CONTENT_W, h, fill=ORANGE)
        size = fit_pt([label + "　" + body], 11.63, h, pt, 13)
        self.text(slide, 0.85, y, 11.63, h, [[(label + "　", {"bold": True}), (body, {})]],
                  pt=size, color=WHITE, anchor="m")

    # -- animation: 指定した図形をクリックごとにフェードイン
    @staticmethod
    def click_fade(slide, shapes):
        P = "http://schemas.openxmlformats.org/presentationml/2006/main"
        ids = [s.shape_id for s in shapes]
        nid = iter(range(3, 1000))
        pars = []
        for spid in ids:
            a, b, c, d, e = (next(nid) for _ in range(5))
            pars.append(f"""
<p:par><p:cTn id="{a}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst>
 <p:par><p:cTn id="{b}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
  <p:par><p:cTn id="{c}" presetID="10" presetClass="entr" presetSubtype="0" fill="hold" nodeType="clickEffect">
   <p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
    <p:set><p:cBhvr><p:cTn id="{d}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>
     <p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr>
     <p:to><p:strVal val="visible"/></p:to></p:set>
    <p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{e}" dur="500"/><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:animEffect>
   </p:childTnLst></p:cTn></p:par>
 </p:childTnLst></p:cTn></p:par>
</p:childTnLst></p:cTn></p:par>""")
        xml = f"""<p:timing xmlns:p="{P}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>
<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>{''.join(pars)}
</p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst><p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>
</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>"""
        el = etree.fromstring(xml)
        sld = slide._element
        ext = sld.find(qn("p:extLst"))
        if ext is not None:
            ext.addprevious(el)
        else:
            sld.append(el)

    # ================================================================ slides
    def build(self):
        d = self.d
        self.cover()
        self.overview()
        self.timing()
        self.three_questions()
        self.basic_form()
        themes = d["themes"]
        for i, th in enumerate(themes, 1):
            if i == 1:
                self.read_slide(i, th, "お手本")
                self.example_q1(i, th)
                self.example_q2(i, th)
                self.example_q3(i, th)
                self.answer_slides(i, th, "お手本", "書く 35分".replace("35", str(self.split()[2])),
                                   "Q3のメモを、そのまま文章にした完成答案")
            else:
                self.read_slide(i, th, "演習")
                self.worksheet1(i, th)
                self.worksheet2(i, th)
                self.model_memo(i, th)
                self.answer_slides(i, th, "答え合わせ", "答え合わせ", th["answer"]["subtitle"])
        self.finish()
        return self.prs

    # -- helpers on exam settings
    def split(self):
        ex = self.d["exam"]
        if ex.get("time_split"):
            return list(ex["time_split"])
        m = ex["minutes"]
        step = 5 if m >= 60 else 1
        r = lambda v: max(step, int(round(v / step)) * step)
        read, plan, review = r(m / 12), r(m / 6), r(m / 6)
        return [read, plan, m - read - plan - review, review]

    def chars_label(self):
        ex = self.d["exam"]
        if ex.get("chars_min"):
            return f"{ex['chars_min']}〜{ex['chars_max']}字"
        return f"{ex['chars_max']}字以内"

    def target_chars(self):
        return int(math.ceil(self.d["exam"]["chars_max"] * 0.9 / 10) * 10)

    def n_themes(self):
        return len(self.d["themes"])

    # -- 1. 表紙
    def cover(self):
        d, meta, ex = self.d, self.d["meta"], self.d["exam"]
        s = self.prs.slides.add_slide(self.blank)
        self.box(s, 0, 0, 0.35, 7.5, fill=GREEN, radius=0)
        head = "　".join(x for x in (meta["university"], meta["faculty"], meta.get("admission", "")) if x)
        self.text(s, 1.0, 1.35, 11.0, 0.5, head, pt=20, color=ORANGE, bold=True)
        self.text(s, 1.0, 1.95, 11.5, 1.2, d["cover"].get("title") or "小論文ワークシート", pt=54, color=INK, bold=True)
        tagline = d["cover"].get("tagline") or f"{ex['minutes']}分で、迷わず、{ex['chars_max']}字を書き切る"
        self.text(s, 1.0, 3.20, 11.5, 0.6, tagline, pt=28, color=GREEN, bold=True)
        n = self.n_themes()
        lead = f"見て学ぶ（お手本）→ やってみる（演習）→ 答え合わせ　を{n}つのテーマで繰り返す" if n > 1 \
            else "見て学ぶ（お手本）→ 自分の答案と見比べる"
        self.text(s, 1.0, 4.05, 11.5, 0.5, lead, pt=18, color=MUTED)
        self.text(s, 1.0, 5.35, 11.5, 0.4, meta["target"], pt=fit_pt([meta["target"]], 11.5, 0.4, 15, 12), color=TEXT)
        self.text(s, 1.0, 5.80, 11.5, 0.4, f"{meta['org']}　学習支援資料", pt=15, color=MUTED)

    # -- 2. 試験の全体像
    def overview(self):
        ex = self.d["exam"]
        s = self.new_slide("INTRODUCTION ｜ 試験を知る", "試験の全体像")
        cards = [(f"{ex['minutes']}分", "小論文の試験時間"),
                 (self.chars_label(), "目安の字数" if ex.get("chars_min") else "字数制限"),
                 (ex["format_label"], ex["format_caption"])]
        for i, (big, small) in enumerate(cards):
            x = 0.60 + i * 4.10
            self.box(s, x, 1.95, 3.85, 2.25, fill=CARD, line=CARD_LINE)
            self.text(s, x, 2.20, 3.85, 1.0, big, pt=fit_pt([big], 3.85, 0.9, 44, 26), color=GREEN, bold=True, align="c")
            self.text(s, x + 0.2, 3.30, 3.45, 0.7, small, pt=fit_pt([small], 3.45, 0.75, 17, 13), color=TEXT, align="c")
        paras = ex["overview"]
        self.text(s, 0.60, 4.50, 12.10, 1.30, paras, pt=fit_pt(paras, 12.1, 1.4, 17, 13, 6), space_after=6)
        self.text(s, 0.60, 5.95, 12.10, 0.5, ex.get("note") or "※ 最新の正式情報は、大学公式の学生募集要項で確認すること。",
                  pt=13, color=MUTED)

    # -- 3. 時間配分
    def timing(self):
        m = self.d["exam"]["minutes"]
        read, plan, write, review = self.split()
        s = self.new_slide("INTRODUCTION ｜ 時間配分", f"{m}分の使い方 ― 考えるのは最初の{read + plan}分",
                           "書きながら考えると、途中で手が止まる。先に「何を書くか」を決めてから、一気に書く。")
        segs = [(read, ORANGE), (plan, ORANGE), (write, GREEN), (review, GREEN)]
        gap = 0.06
        unit = (CONTENT_W - 3 * gap) / m
        widths = [max(0.8, v * unit) for v, _ in segs]
        scale = (CONTENT_W - 3 * gap - 0) / sum(widths)
        widths = [w * scale for w in widths]
        x = LEFT
        for (v, col), w in zip(segs, widths):
            self.box(s, x, 2.20, w, 0.70, fill=col)
            self.text(s, x - 0.03, 2.20, w + 0.06, 0.70, f"{v}分", pt=16 if w < 1.2 else 20, color=WHITE, bold=True,
                      align="c", anchor="m")
            x += w + gap
        steps = [("最初の", read, "読む", ["問題文と資料を読み、", "要点をつかむ"], ORANGE),
                 ("次の", plan, "構成する", ["骨組みと各段落の中身を", "箇条書きでメモ"], ORANGE),
                 ("続く", write, "書く", ["メモに沿って", f"一気に{self.chars_label()}を書き切る"], GREEN),
                 ("最後の", review, "見直す", ["誤字脱字・文末の統一・", "論理のつながりを確認"], GREEN)]
        for i, (pre, v, verb, desc, col) in enumerate(steps):
            x = 0.60 + i * 3.08
            self.box(s, x, 3.20, 2.90, 2.75, fill=CARD, line=CARD_LINE)
            self.text(s, x + 0.2, 3.35, 2.6, 0.4, f"{pre}{v}分", pt=16, color=col, bold=True)
            self.text(s, x + 0.2, 3.80, 2.6, 0.7, verb, pt=30, color=INK, bold=True)
            self.text(s, x + 0.2, 4.65, 2.6, 1.2, desc, pt=fit_pt(desc, 2.6, 1.2, 16, 12))
        self.bar(s, "POINT", f"ワークシート（次のページ）は「構成する{plan}分」でやることそのもの。")

    # -- 4. 3つの問い
    def three_questions(self):
        s = self.new_slide("INTRODUCTION ｜ ワークシートの使い方", "構成メモは、3つの問いで作る",
                           "どのテーマでも、この順番で考えて、ノート（解答用紙）に書き出す。")
        qs = [("Q1", "この問題・資料の「要点」は何か",
               ["・何を問われているか（説明？意見？具体策？）", "・資料から必ず使う事実・数字は何か", "・一言でいうと、何について書く問題か"]),
              ("Q2", "全体を何段落で書くか",
               ["・序論・本論・結論をそれぞれ何段落にするか", "・問いが2つなら本論も2つ、が基本", "・各段落のおおよその字数を決める"]),
              ("Q3", "各段落に書く内容は何か",
               ["・段落ごとに、書く中身を箇条書きにする", "・文章にしなくてよい（単語・短いメモで十分）", "・このメモがそのまま答案の設計図になる"])]
        for i, (q, head, items) in enumerate(qs):
            x = 0.60 + i * 4.10
            self.box(s, x, 2.10, 3.85, 3.80, fill=CARD, line=CARD_LINE)
            self.pill(s, x + 0.25, 2.30, 0.90, 0.55, q, fill=ORANGE, pt=20, radius=6000)
            self.text(s, x + 0.25, 3.00, 3.45, 0.9, head, pt=19, color=INK, bold=True)
            self.text(s, x + 0.25, 3.95, 3.45, 1.9, items, pt=15, space_after=4)
        n = self.n_themes()
        flow = "テーマ1はお手本で学ぶ → " + (f"テーマ2〜{n}は自分で埋める → 模範メモ・模範解答で答え合わせ" if n > 2
                                         else "テーマ2は自分で埋める → 模範メモ・模範解答で答え合わせ" if n == 2
                                         else "同じ手順で過去問・類題に取り組む")
        self.bar(s, "流れ", flow)

    # -- 5. 基本フォーム
    def basic_form(self):
        b = self.d["basics"]
        s = self.new_slide("INTRODUCTION ｜ 書き方の基本", "答案の基本フォーム")
        cols = [("序論　15〜20%", ["自分の意見・結論を", "先に示す"]),
                ("本論　60〜70%", ["理由・根拠・具体例を", "2〜3点で展開する"]),
                ("結論　15〜20%", ["主張を確認し、", "大学での学びにつなげる"])]
        for i, (head, desc) in enumerate(cols):
            x = 0.60 + i * 4.10
            self.pill(s, x, 1.85, 3.85, 0.60, head, fill=GREEN, pt=20, radius=6000)
            self.box(s, x, 2.50, 3.85, 1.30, fill=CARD, line=CARD_LINE)
            self.text(s, x + 0.2, 2.60, 3.45, 1.10, desc, pt=18, align="c", anchor="m")
        rules = b.get("style_rules") or ["「だ・である」調で統一する", "一文を短く、一文一意", "話し言葉（〜だと思う、〜なので）を避ける"]
        for x, w, head, items in ((0.60, 6.00, "文体のルール", rules),
                                  (6.75, 5.98, "大学が見ていること", b["evaluation_points"])):
            self.box(s, x, 4.05, w, 2.00, line=CARD_LINE)
            body = ["・" + t for t in items]
            pt = fit_pt(body, w - 0.4, 1.4, 16, 12, 3)
            self.text(s, x + 0.25, 4.15, w - 0.4, 1.85,
                      [[(head, {"pt": 18, "bold": True, "color": GREEN})]] + body, pt=pt, space_after=3)
        ex = self.d["exam"]
        self.bar(s, "字数", f"{self.chars_label()}なら、最低でも{self.target_chars()}字（9割）を目標に書き切る。")

    # -- テーマ: 読む
    def read_slide(self, i, th, role):
        read = self.split()[0]
        s = self.new_slide(f"THEME {i} ｜ {role}", f"テーマ{i}｜{th['title']}",
                           f"{th['source']}／{self.chars_label()}", badge=f"読む {read}分")
        self.box(s, LEFT, 1.95, 7.30, 4.15, fill=CARD, line=CARD_LINE)
        self.text(s, 0.85, 2.05, 6.90, 0.40, "問題（要旨）", pt=16, color=GREEN, bold=True)
        mat = th.get("material") or {"type": "none"}
        mt = mat.get("type", "none")
        q = th["question"]
        q_room = 3.55 if mt == "none" else 1.75
        qpt = fit_pt([q], 6.9, q_room, 17, 13)
        qh = est_height([q], 6.9, qpt)
        self.text(s, 0.85, 2.45, 6.90, qh, q, pt=qpt)
        y = 2.45 + qh + 0.12
        bottom = 6.00
        if mt == "table":
            cols, rows = mat["columns"], mat["rows"]
            if mat.get("heading"):
                self.text(s, 0.85, y, 6.9, 0.4, mat["heading"], pt=15, color=GREEN, bold=True)
                y += 0.42
            rh = min(0.475, (bottom - y) / (len(rows) + 1))
            tpt = 15 if rh >= 0.45 else 13 if rh >= 0.36 else 11
            gt = s.shapes.add_table(len(rows) + 1, len(cols), Inches(0.85), Inches(y), Inches(6.80),
                                    Inches(rh * (len(rows) + 1)))
            tbl = gt.table
            for c in range(len(cols)):
                tbl.columns[c].width = Inches(6.80 / len(cols))
            for r in range(len(rows) + 1):
                tbl.rows[r].height = Inches(rh)
                vals = cols if r == 0 else rows[r - 1]
                for c in range(len(cols)):
                    cell = tbl.cell(r, c)
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = rgb(GREEN if r == 0 else WHITE)
                    cell.margin_top = cell.margin_bottom = Inches(0.02)
                    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                    p = cell.text_frame.paragraphs[0]
                    p.alignment = PP_ALIGN.CENTER
                    self._run(p, str(vals[c]) if c < len(vals) else "", tpt, r == 0, WHITE if r == 0 else TEXT)
            tblPr = tbl._tbl.tblPr
            tblPr.set("firstRow", "1"); tblPr.set("bandRow", "0")
            if mat.get("note"):
                self.text(s, 0.85, y + rh * (len(rows) + 1) + 0.02, 6.9, 0.3, mat["note"], pt=11, color=MUTED)
        elif mt == "chips":
            items = mat["items"]
            self.text(s, 0.85, y, 6.9, 0.4, mat.get("heading") or "資料のキーワード", pt=15, color=GREEN, bold=True)
            y += 0.48
            longest = max(text_units(t) for t in items)
            ncol = 3 if longest <= 8 and len(items) > 4 else 2
            cw = (6.75 - (ncol - 1) * 0.15) / ncol
            nrow = math.ceil(len(items) / ncol)
            ch = min(0.52, (bottom - y - (nrow - 1) * 0.1) / nrow)
            for k, t in enumerate(items):
                r, c = divmod(k, ncol)
                self.pill(s, 0.85 + c * (cw + 0.15), y + r * (ch + 0.1), cw, ch, t, fill=WHITE, color=GREEN,
                          pt=fit_pt([t], cw, ch, 15, 11), line=GREEN, radius=30000)
        elif mt in ("conditions", "text"):
            head = mat.get("heading") or ("条件" if mt == "conditions" else "資料（要旨）")
            body = ["・" + t for t in mat["items"]] if mt == "conditions" else [mat["text"]]
            pt = fit_pt(body, 6.9, bottom - y - 0.4, 15, 11, 2)
            self.text(s, 0.85, y, 6.90, bottom - y,
                      [[(head, {"bold": True, "color": GREEN})]] + body, pt=pt, space_after=2)
        # 線を引くところ
        self.box(s, 8.10, 1.95, 4.63, 4.15, fill=PEACH)
        self.text(s, 8.35, 2.05, 4.20, 0.40, "読むときに線を引くところ", pt=16, color=ORANGE, bold=True)
        pts = [CIRCLED[k] + " " + t for k, t in enumerate(th["underline_points"])]
        self.text(s, 8.35, 2.50, 4.20, 3.50, pts, pt=fit_pt(pts, 4.2, 3.5, 15, 11, 6), space_after=6)
        self.bar(s, f"最初の{read}分", "問題文と資料を読み、「何を問われているか」に線を引く。まだ書き始めない。")

    # -- テーマ1 お手本: Q1
    def example_q1(self, i, th):
        q1 = th["q1"]
        plan = self.split()[1]
        s = self.new_slide(f"THEME {i} ｜ お手本", "お手本｜Q1　要点をつかむ",
                           "クリックするたびに、考える順番どおりに表示されます。", badge=f"構成 {plan}分")
        rows = [(a["label"], a["text"], False) for a in q1["asked"]]
        rows += [("資料のポイント", q1["material_point"], False), ("一言でいうと", q1["one_line"], True)]
        top, bottom = 1.95, 6.00
        rh = min(0.88, (bottom - top - (len(rows) - 1) * 0.12) / len(rows))
        groups = []
        for k, (label, body, hl) in enumerate(rows):
            y = top + k * (rh + 0.12)
            parts = [self.box(s, LEFT, y, CONTENT_W, rh, fill=PEACH if hl else CARD,
                              line=ORANGE if hl else CARD_LINE, line_w=1.5 if hl else 1.0),
                     self.pill(s, 0.80, y + (rh - 0.48) / 2, 2.60, 0.48, label, fill=ORANGE if hl else GREEN, pt=15,
                               radius=6000)]
            pt = fit_pt([body], 8.9, rh, 19 if hl else 18, 12)
            parts.append(self.text(s, 3.60, y, 8.90, rh, body, pt=pt, bold=hl, anchor="m"))
            groups.append(self.group(s, parts))
        self.bar(s, "コツ", th.get("tip") or "設問の「〜について説明」「〜について述べよ」を数えると、答える項目の数が分かる。",
                 y=6.10)
        self.click_fade(s, groups)

    # -- テーマ1 お手本: Q2
    def example_q2(self, i, th):
        q2 = th["q2"]
        plan = self.split()[1]
        s = self.new_slide(f"THEME {i} ｜ お手本", "お手本｜Q2　段落の数を決める", q2["lead"], badge=f"構成 {plan}分")
        ps = q2["paragraphs"]
        n = len(ps)
        gap = 0.15
        cw = (CONTENT_W - (n - 1) * gap) / n
        groups = []
        for k, p in enumerate(ps):
            x = LEFT + k * (cw + gap)
            parts = [self.box(s, x, 2.20, cw, 2.60, fill=CARD, line=CARD_LINE),
                     self.pill(s, x + 0.15, 2.35, cw - 0.30, 0.55, p["label"], fill=GREEN, pt=18, radius=6000),
                     self.text(s, x + 0.10, 3.05, cw - 0.20, 0.50, p["chars"], pt=fit_pt([p["chars"]], cw - 0.2, 0.5, 17, 12),
                               color=ORANGE, bold=True, align="c"),
                     self.text(s, x + 0.10, 3.60, cw - 0.20, 1.10, p["role"],
                               pt=fit_pt([p["role"]], cw - 0.2, 1.1, 16, 11), align="c")]
            groups.append(self.group(s, parts))
        dec = [self.box(s, LEFT, 5.00, CONTENT_W, 0.95, fill=PEACH, line=ORANGE, line_w=1.5),
               self.text(s, 0.85, 5.00, 11.63, 0.95, [[("決定：", {"color": ORANGE}), (q2["decision"], {})]],
                         pt=fit_pt(["決定：" + q2["decision"]], 11.63, 0.9, 20, 14), bold=True, anchor="m")]
        groups.append(self.group(s, dec))
        self.bar(s, q2.get("why_title") or "なぜこの構成？", q2["why"], y=6.10)
        self.click_fade(s, groups)

    # -- 段落ラベル付きの行（Q3 メモなど）
    def label_rows(self, s, rows, top, bottom, x=LEFT, w=CONTENT_W, label_w=1.15, pt=16, fill=CARD, gap=0.07,
                   max_h=1.1):
        n = len(rows)
        rh = min(max_h, (bottom - top - (n - 1) * gap) / n)
        tx = x + label_w + 0.30
        tw = x + w - tx - 0.12
        size = min(fit_pt([t], tw, rh - 0.06, pt, 11) for _, t in rows)
        shapes = []
        for k, (label, body) in enumerate(rows):
            y = top + k * (rh + gap)
            parts = [self.box(s, x, y, w, rh, fill=fill, line=CARD_LINE),
                     self.pill(s, x + 0.12, y + 0.08, label_w, 0.40, label, fill=GREEN, pt=14, radius=30000),
                     self.text(s, tx, y + 0.04, tw, rh - 0.06, body, pt=size)]
            shapes.append(parts)
        return shapes

    # -- テーマ1 お手本: Q3
    def example_q3(self, i, th):
        plan = self.split()[1]
        read = self.split()[0]
        s = self.new_slide(f"THEME {i} ｜ お手本", "お手本｜Q3　各段落を箇条書きで肉付けする",
                           "文章ではなく、短いメモでよい。このメモがそのまま答案になる。", badge=f"構成 {plan}分")
        rows = [(m["label"], m["text"]) for m in th["q3"]["memo"]]
        parts = self.label_rows(s, rows, 1.95, 6.13)
        groups = [self.group(s, p) for p in parts]
        write = self.split()[2]
        self.bar(s, f"ここまで{read + plan}分", f"メモが完成したら、あとは{write}分で書き写すように文章にするだけ。", y=6.20, h=0.60)
        self.click_fade(s, groups)

    # -- 模範解答（長い答案は2枚に分ける）
    def answer_slides(self, i, th, role, badge, sub):
        ans = th["answer"]
        paras = ans["paragraphs"]
        counts = [count_chars(p["text"]) for p in paras]
        total = sum(counts)
        # ラベルに数字が入る（設問1 など）ときは、字数と混ざらないよう「：」で区切る
        detail = "・".join(f"{p['label'].replace(' ', '')}{'：' if any(ch in '0123456789' for ch in p['label']) else ''}{c}"
                          for p, c in zip(paras, counts))
        tw = 10.58
        box_h = 4.42

        def layout(chunk, pt):
            # 実測の行送り（約1.3em）＋上下余白で各段落の高さを決める
            hs = [max(0.50, est_lines(p["text"], tw, pt) * pt * 1.3 / 72 + 0.06) for p in chunk]
            return hs, sum(hs) + 0.07 * (len(chunk) - 1) + 0.24

        # 1枚に 12pt 以上で収まらなければ、段落の切れ目で2枚に分ける
        pt = 13.5
        while pt > 12 and layout(paras, pt)[1] > box_h:
            pt -= 0.5
        if layout(paras, pt)[1] <= box_h:
            chunks = [paras]
        else:
            pt = 13.5
            cut = min(range(1, len(paras)),
                      key=lambda c: max(layout(paras[:c], pt)[1], layout(paras[c:], pt)[1]))
            chunks = [paras[:cut], paras[cut:]]
            while pt > 11 and max(layout(c, pt)[1] for c in chunks) > box_h:
                pt -= 0.5
        for ci, chunk in enumerate(chunks):
            title = f"テーマ{i}｜模範解答" + (f"（{ci + 1}/{len(chunks)}）" if len(chunks) > 1 else "")
            s = self.new_slide(f"THEME {i} ｜ {role}", title, sub, badge=badge)
            self.text(s, 0.60, 1.48, 12.10, 0.40, f"合計{total}字（{detail}）", pt=13, color=MUTED, align="r")
            self.box(s, LEFT, 1.90, CONTENT_W, box_h, fill=CARD, line=CARD_LINE)
            hs, _ = layout(chunk, pt)
            y = 2.02
            for p, h in zip(chunk, hs):
                self.pill(s, 0.77, y + 0.04, 1.00, 0.40, p["label"], fill=GREEN, pt=14, radius=30000)
                self.text(s, 1.95, y, tw, h, p["text"], pt=pt)
                y += h + 0.07
            pts = "　".join("✓ " + t for t in ans["points"])
            self.bar(s, "ポイント", pts, y=6.42, h=0.50, pt=15)

    # -- 演習: ワークシート①
    def worksheet1(self, i, th):
        plan = self.split()[1]
        s = self.new_slide(f"THEME {i} ｜ 演習", f"テーマ{i}｜ワークシート ①",
                           f"{th['title']}　― ノート（解答用紙）に書き出そう", badge=f"構成 {plan}分")
        for x, q, head, hints in ((0.60, "Q1", "この問題・資料の「要点」は何か", th["hints"]["q1"]),
                                  (6.75, "Q2", "全体を何段落で書くか", th["hints"]["q2"])):
            self.pill(s, x, 2.00, 0.85, 0.50, q, fill=ORANGE, pt=18, radius=6000)
            self.text(s, x + 1.0, 1.98, 5.0, 0.55, head, pt=20, color=INK, bold=True, anchor="m")
            self.box(s, x, 2.65, 5.98, 2.05, fill=CARD)
            body = ["・" + h for h in hints]
            pt = fit_pt(body, 5.6, 1.55, 15, 11, 3)
            self.text(s, x + 0.2, 2.72, 5.6, 1.95, [[("ヒント", {"bold": True, "color": GREEN, "pt": 15})]] + body,
                      pt=pt, space_after=3)
            self.box(s, x, 4.85, 5.98, 1.25, line=ORANGE, line_w=1.75, dash=True)
            self.text(s, x + 0.2, 4.93, 5.6, 0.4, "あなたの答え（ノートに書く）", pt=14, color=ORANGE, bold=True)
        self.bar(s, "目安", "Q1・Q2で3〜4分。迷ったら、問われていることの数＝本論の段落数で決める。")

    # -- 演習: ワークシート②
    def worksheet2(self, i, th):
        plan = self.split()[1]
        write, review = self.split()[2], self.split()[3]
        s = self.new_slide(f"THEME {i} ｜ 演習", f"テーマ{i}｜ワークシート ②",
                           "Q3　各段落に書く内容を、箇条書きで（文章にしなくてよい）", badge=f"構成 {plan}分")
        rows = th["q3"]["prompts"]
        n = len(rows)
        top, bottom, gap = 1.95, 6.15, 0.08
        rh = min(0.75, (bottom - top - (n - 1) * gap) / n)
        size = min(fit_pt([r["text"]], 4.6, rh, 14, 11) for r in rows)
        for k, r in enumerate(rows):
            y = top + k * (rh + gap)
            self.box(s, LEFT, y, CONTENT_W, rh, line=CARD_LINE)
            self.pill(s, 0.75, y + (rh - 0.42) / 2, 1.45, 0.42, r["label"], fill=GREEN, pt=14, radius=30000)
            self.text(s, 2.35, y + 0.02, 4.60, rh - 0.04, r["text"], pt=size, color=MUTED, anchor="m")
            self.box(s, 7.05, y + 0.07, 5.55, rh - 0.14, line=ORANGE, line_w=1.25, dash=True)
        self.bar(s, "メモができたら", f"続く{write}分で、このメモどおりに一気に書く → 最後の{review}分で見直す。", y=6.30, h=0.58)

    # -- 演習: 模範メモ
    def model_memo(self, i, th):
        s = self.new_slide(f"THEME {i} ｜ 答え合わせ", f"テーマ{i}｜模範メモ",
                           f"{th['title']}　― 自分のメモと見比べよう", badge="答え合わせ")
        self.box(s, LEFT, 1.95, 5.00, 4.15, fill=CARD, line=CARD_LINE)
        q1 = ["・" + t for t in th["q1"]["memo"]]
        q2 = ["・" + t for t in th["q2"]["memo"]]
        pt = fit_pt(q1 + q2, 4.6, 4.0 - 0.9, 15, 11, 3)
        self.text(s, 0.80, 2.05, 4.60, 4.00,
                  [[("Q1　要点", {"pt": 17, "bold": True, "color": ORANGE})]] + q1 +
                  [[("Q2　段落構成", {"pt": 17, "bold": True, "color": ORANGE})]] + q2, pt=pt, space_after=3)
        self.text(s, 5.85, 1.95, 6.90, 0.40, "Q3　各段落のメモ", pt=17, color=ORANGE, bold=True)
        rows = [(m["label"], m["text"]) for m in th["q3"]["memo"]]
        self.label_rows(s, rows, 2.40, 6.05, x=5.85, w=6.88, label_w=1.05, pt=14, fill=WHITE, gap=0.08, max_h=0.9)
        self.bar(s, "チェック", "問われていることに全部答えているか。段落の役割が重なっていないか。")

    # -- 最後: 見直し
    def finish(self):
        f = self.d["final"]
        read, plan, write, review = self.split()
        s = self.new_slide("FINISH ｜ 最後の見直し", "見直しチェックリストとまとめ", badge=f"見直す {review}分")
        self.box(s, LEFT, 1.75, 6.60, 4.35, fill=CARD, line=CARD_LINE)
        must = ["✓ " + t for t in (f.get("review_points") or
                                   ["誤字脱字・送りがな・数字の書き間違い", "文末が「だ・である」で統一されているか",
                                    "論理のつながり（主張 → 理由 → 結論がずれていないか）"])]
        extra = ["✓ " + t for t in f["extra_checks"]] + [f"✓ {self.target_chars()}字（9割）以上書けているか"]
        pt = fit_pt(must + extra, 6.2, 3.3, 16, 12, 4)
        self.text(s, 0.85, 1.85, 6.20, 0.40, f"最後の{review}分で必ず見る{len(must)}点", pt=18, color=ORANGE, bold=True)
        self.text(s, 0.85, 2.35, 6.20, 1.50, must, pt=pt, space_after=4)
        self.text(s, 0.85, 3.90, 6.20, 0.40, "あわせて確認", pt=18, color=GREEN, bold=True)
        self.text(s, 0.85, 4.35, 6.20, 1.70, extra, pt=pt, space_after=4)
        for k, (v, verb, col) in enumerate(((read, "読む", ORANGE), (plan, "構成する", ORANGE),
                                            (write, "書く", GREEN), (review, "見直す", GREEN))):
            y = 1.75 + k * 1.12
            self.pill(s, 7.50, y, 1.50, 0.95, f"{v}分", fill=col, pt=22, radius=6000)
            self.text(s, 9.20, y, 3.50, 0.95, verb, pt=26, color=INK, bold=True, anchor="m")
        self.bar(s, "まとめ", f"考えるのは最初の{read + plan}分。3つの問いで、初見の問題も同じ手順で書き切れる。", y=6.30, h=0.60)


def render(data, out_path):
    from schema import validate
    validate(data)
    prs = Deck(data).build()
    prs.save(out_path)
    return out_path


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit("usage: python3 tools/render.py content.json out.pptx")
    with open(sys.argv[1], encoding="utf-8") as fh:
        render(json.load(fh), sys.argv[2])
    print(sys.argv[2])
